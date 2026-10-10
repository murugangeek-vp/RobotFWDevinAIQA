"""Write the migration evidence artifacts.

migration_summary.xlsx — Run Info / Stage Results / Diffs & Violations /
                         Schema Errors / Insights sheets (the audit
                         workbook a reviewer opens next to log.html)
migration_report.md    — the same content as a readable QA report, plus
                         auto-generated Insights and Improvement
                         Suggestions for the review meeting.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_FILL = {"PASS": "C6EFCE", "WARN": "FFEB9C", "FAIL": "FFC7CE", "ERROR": "FFC7CE", "SKIP": "DDDDDD"}


def _bold_header(ws) -> None:
    for c in ws[1]:
        c.font = Font(bold=True)


def _autofit(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, 80)


def build_insights(stage_results, recon_result, schema_errors, rule_failures, meta) -> dict:
    """Deterministic insight + improvement-suggestion generation — evidence
    first, then actionable follow-ups a reviewer can assign."""
    insights, suggestions = [], []
    stages = {s["stage"]: s for s in stage_results}
    failed = [s for s in stage_results if str(s["status"]).upper() not in ("PASS", "SKIP")]

    insights.append(f"Stages executed: {len(stage_results)}; failed: {len(failed)}")
    if recon_result:
        insights.append(
            f"Rows compared — source: {recon_result.source_count}, "
            f"target: {recon_result.target_count}"
        )
        if recon_result.missing_in_target:
            n = len(recon_result.missing_in_target)
            insights.append(f"{n} source key(s) absent from target")
            suggestions.append(
                f"Investigate the load layer: {n} source records never reached the "
                "target (samples in Diffs & Violations). Check ETL filters and rejects."
            )
        if recon_result.extra_in_target:
            n = len(recon_result.extra_in_target)
            insights.append(f"{n} extra key(s) present in target")
            suggestions.append(
                f"Review migration scope: {n} target rows have no source counterpart — "
                "leaked out-of-scope rows or stale data. Tighten the WHERE clause or "
                "post-migration cleanup."
            )
        if recon_result.diffs:
            cols = sorted({d.column for d in recon_result.diffs})
            insights.append(
                f"{len(recon_result.diffs)} column value diff(s) across {len(cols)} column(s)"
            )
            suggestions.append(
                f"Transformation drift on {cols}: re-check the mapping workbook "
                "Transformation expressions and the ETL implementation."
            )
        if recon_result.mismatch_count == 0:
            insights.append("All compared rows reconcile exactly — zero mismatches.")
            suggestions.append(
                "Reconciliation is clean. Schedule periodic verify-only runs "
                "(RECON_SKIP_LOAD) post-cutover to catch drift in production."
            )
    if schema_errors:
        insights.append(f"Schema validation produced {len(schema_errors)} error(s)")
        suggestions.append(
            "Align the target DDL with the contract: schema drift detected "
            "(see Schema Errors sheet)."
        )
    if rule_failures:
        n = sum(f.failing_rows for f in rule_failures)
        insights.append(f"Source data quality: {n} rule violation(s)")
        suggestions.append(
            "Cleanse source data before migration (see DQ violations) or extend "
            "the mapping to handle them explicitly."
        )
    if recon_result and recon_result.source_count > 200000:
        suggestions.append(
            "Table volume justifies recon.mode: hash in the mapping workbook — "
            "key->rowhash comparison is the fast path for million-row tables."
        )
    unrun = [s for s in ("schema", "recon", "dq") if s not in stages]
    if unrun:
        insights.append(f"Stages not exercised this run: {', '.join(unrun)}")
    if not suggestions:
        suggestions.append("No issues found — no follow-up actions.")
    return {"insights": insights, "suggestions": suggestions}


def _verdict(stage_results, recon_result, schema_errors, rule_failures) -> str:
    if any(str(s["status"]).upper() not in ("PASS", "SKIP") for s in stage_results):
        return "FAIL"
    if schema_errors or rule_failures:
        return "FAIL"
    if recon_result and recon_result.mismatch_count:
        return "FAIL"
    return "PASS"


def write_migration_report(
    stage_results: list,
    recon_result,
    schema_errors: list,
    rule_failures: list,
    meta: dict,
    out_dir: "str | Path" = "results",
) -> dict:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    derived = build_insights(stage_results, recon_result, schema_errors, rule_failures, meta)
    verdict = _verdict(stage_results, recon_result, schema_errors, rule_failures)
    xlsx = out / "migration_summary.xlsx"
    md = out / "migration_report.md"

    wb = Workbook()
    info = wb.active
    info.title = "Run Info"
    info.append(["Field", "Value"])
    _bold_header(info)
    for field, value in [
        ("Run Timestamp", meta.get("timestamp", "")),
        ("Environment", meta.get("environment", "")),
        ("Contract / Mapping", meta.get("contract", "")),
        ("Mapping Workbook", meta.get("mapping_workbook", "")),
        ("Executed By", meta.get("executed_by", "")),
        ("Git Commit", meta.get("git_commit", "")),
        ("Stages Run", len(stage_results)),
        ("Verdict", verdict),
    ]:
        info.append([field, value])
        info.cell(row=info.max_row, column=1).font = Font(bold=True)
    info.cell(row=info.max_row, column=2).fill = PatternFill("solid", fgColor=_FILL[verdict])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 70

    ws = wb.create_sheet("Stage Results")
    ws.append(["Mapping", "Stage", "Status", "Detail", "Timestamp"])
    _bold_header(ws)
    for s in stage_results:
        ws.append(
            [s.get("mapping", ""), s["stage"], s["status"], s.get("detail", ""), s["timestamp"]]
        )
        fill = _FILL.get(str(s["status"]).upper())
        if fill:
            ws.cell(row=ws.max_row, column=3).fill = PatternFill("solid", fgColor=fill)
    _autofit(ws)

    ws = wb.create_sheet("Diffs & Violations")
    ws.append(["Kind", "Key", "Column", "Expected", "Actual"])
    _bold_header(ws)
    if recon_result:
        for k in recon_result.missing_in_target[:500]:
            ws.append(["missing_in_target", str(k), "", "", ""])
        for k in recon_result.extra_in_target[:500]:
            ws.append(["extra_in_target", str(k), "", "", ""])
        for d in recon_result.diffs[:500]:
            ws.append(["value_diff", str(d.key), d.column, str(d.expected), str(d.actual)])
    _autofit(ws)

    ws = wb.create_sheet("Schema Errors")
    ws.append(["Error"])
    _bold_header(ws)
    for e in schema_errors or []:
        ws.append([e])
    for f in rule_failures or []:
        ws.append([f"DQ {f.rule_id}: {f.failing_rows} rows, samples {f.samples}"])
    _autofit(ws)

    ws = wb.create_sheet("Insights")
    ws.append(["Type", "Detail"])
    _bold_header(ws)
    for i in derived["insights"]:
        ws.append(["insight", i])
    for s in derived["suggestions"]:
        ws.append(["suggestion", s])
    _autofit(ws)

    wb.save(xlsx)

    lines = [
        "# Migration Reconciliation Report",
        "",
        f"- **Run:** {meta.get('timestamp', '')} | **Environment:** {meta.get('environment', '')}",
        f"- **Mapping:** {meta.get('contract', '')} | **Workbook:** {meta.get('mapping_workbook', '')}",
        f"- **Executed by:** {meta.get('executed_by', '')} | **Commit:** {meta.get('git_commit', '')}",
        f"- **Verdict:** {verdict}",
        "",
        "## Stage Results",
        "",
        "| Mapping | Stage | Status | Detail |",
        "|---|---|---|---|",
    ]
    for s in stage_results:
        lines.append(
            f"| {s.get('mapping', '')} | {s['stage']} | {s['status']} | {s.get('detail', '')} |"
        )
    lines += ["", "## Reconciliation", ""]
    if recon_result:
        lines += [
            f"- Source rows: {recon_result.source_count} | Target rows: {recon_result.target_count}",
            f"- Missing in target: {len(recon_result.missing_in_target)}",
            f"- Extra in target: {len(recon_result.extra_in_target)}",
            f"- Column diffs: {len(recon_result.diffs)}",
        ]
        for d in recon_result.diffs[:20]:
            lines.append(
                f"  - key `{d.key}` col `{d.column}`: expected {d.expected}, got {d.actual}"
            )
    else:
        lines.append("- No comparison ran this session.")
    if schema_errors:
        lines += ["", "## Schema Errors", ""] + [f"- {e}" for e in schema_errors]
    if rule_failures:
        lines += ["", "## DQ Rule Failures", ""] + [
            f"- `{f.rule_id}`: {f.failing_rows} rows (samples {f.samples})" for f in rule_failures
        ]
    lines += (
        ["", "## Insights", ""]
        + [f"- {i}" for i in derived["insights"]]
        + ["", "## Improvement Suggestions", ""]
        + [f"- {s}" for s in derived["suggestions"]]
    )
    md.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return {"xlsx": str(xlsx), "md": str(md), "verdict": verdict}


def write_evidence(results, meta, out_dir):
    import json

    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    verdict = (
        "NOT_RUN"
        if not results
        else ("PASS" if all(r["status"] == "PASS" for r in results) else "FAIL")
    )
    required = {"source", "dq", "schema", "recon", "transform"}
    coverage = {
        name: sorted(required - {r["stage"] for r in results if r["mapping"] == name})
        for name in meta.get("selected", [])
    }
    complete = bool(coverage) and not any(coverage.values())
    cutover = {
        "snapshot_id": meta.get("snapshot_id", ""),
        "source_snapshot": meta.get("source_snapshot", ""),
        "snapshot_enforced": bool(meta.get("snapshot_enforced")),
        "post_snapshot_arrivals": {
            r["mapping"]: r["post_snapshot_arrivals"]
            for r in results
            if r.get("stage") == "recon" and r.get("post_snapshot_arrivals") is not None
        },
        "target_asof": {r["mapping"]: r["target_asof"] for r in results if r.get("target_asof")},
    }
    summary = {
        "verdict": verdict,
        "coverage": "COMPLETE" if complete else "PARTIAL",
        "not_run": coverage,
        "meta": meta,
        "cutover": cutover,
        "results": results,
        "production_signoff": False,
    }
    (out / "run_summary.json").write_text(
        json.dumps(summary, indent=2, default=str), encoding="utf-8"
    )
    wb = Workbook()
    wb.active.title = "Run Info"
    sheets = {
        "Run Info": [
            ["Field", "Value"],
            ["Verdict", verdict],
            ["Coverage", summary["coverage"]],
            ["Production Signoff", "NOT GRANTED: real account and snapshot validation required"],
        ],
        "Stage Results": [
            [
                "Mapping",
                "Stage",
                "Status",
                "Source Rows",
                "Target Rows",
                "Missing",
                "Extra",
                "Column Diffs",
                "Seconds",
                "Detail",
            ]
        ],
        "Metadata": [["Mapping", "Stage", "Source", "Target", "Observed Metadata"]],
        "Diffs & Violations": [
            ["Mapping", "Stage", "Evidence (redacted samples; totals in Stage Results)"]
        ],
        "Insights": [["Type", "Detail"]],
    }
    for key, value in meta.items():
        sheets["Run Info"].append(
            [key, json.dumps(value, default=str) if isinstance(value, dict | list) else str(value)]
        )
    for key, value in cutover.items():
        sheets["Run Info"].append([f"cutover.{key}", json.dumps(value, default=str)])
    for row in results:
        sheets["Stage Results"].append(
            [
                row.get(k, "")
                for k in (
                    "mapping",
                    "stage",
                    "status",
                    "source_count",
                    "target_count",
                    "missing",
                    "extra",
                    "diffs",
                    "seconds",
                    "detail",
                )
            ]
        )
        observed = {
            k: row[k]
            for k in (
                "source_columns",
                "source_metadata",
                "target_metadata",
                "field_mappings",
                "unmapped_source_tables",
                "unmapped_target_tables",
            )
            if k in row
        }
        if observed:
            sheets["Metadata"].append(
                [
                    row["mapping"],
                    row["stage"],
                    row.get("source", ""),
                    row.get("target", ""),
                    json.dumps(observed, default=str),
                ]
            )
        for key in ("dq_failures", "schema_errors", "diff_samples"):
            if row.get(key):
                sheets["Diffs & Violations"].append(
                    [row["mapping"], row["stage"], json.dumps(row[key], default=str)]
                )
    insights = [
        "Results cover only the selected mappings and executed stages; no unmapped table is implicitly certified.",
        "Millions-of-rows production scale and real Snowflake role enforcement require separate environment validation.",
        "Comparison uses bounded fetch batches plus private temporary disk sorting; reserve disk proportional to projected rows.",
        "Client-side hash mode does not reduce database network transfer and is not SQL hash pushdown.",
        "Use an agreed immutable migration snapshot on both sides; separate live systems have no shared transaction snapshot.",
        "Row counts alone cannot validate account balances. Review amount sums and referential integrity as additional banking controls.",
        "Negative tests alter streamed copies, not database tables; they prove detection, not ETL recovery.",
    ]
    if meta.get("emulated_target"):
        insights.insert(
            0,
            "EMULATED TARGET: fakesnow is not a real Santander or Snowflake deployment. All demo data is synthetic.",
        )
    if meta.get("snapshot_enforced"):
        insights.insert(
            0,
            f"Source snapshot ENFORCED ({meta.get('source_snapshot')}); "
            f"cut id: {meta.get('snapshot_id') or 'exported pg snapshot'}.",
        )
    else:
        insights.insert(
            0,
            "WARNING: source snapshot was recorded but NOT enforced — rows may "
            "have changed between stage reads.",
        )
    if any(cutover["post_snapshot_arrivals"].values()):
        insights.append(
            f"Post-snapshot arrivals detected: {cutover['post_snapshot_arrivals']} "
            "— these rows are outside the validated cut and need a delta migration."
        )
    if verdict == "FAIL":
        insights.append(
            "Resolve every failed/error stage before sign-off: align DDL, repair rejected source rows, and review mapping transformations."
        )
    for name, missing in coverage.items():
        if missing:
            insights.append(f"{name}: stages not run: {', '.join(missing)}")
    sheets["Insights"].extend([["Insight / follow-up", item] for item in insights])
    for title, rows in sheets.items():
        ws = wb[title] if title in wb.sheetnames else wb.create_sheet(title)
        for row in rows:
            ws.append(row)
        for cells in ws:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        _bold_header(ws)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        _autofit(ws)
    wb.save(out / "migration_summary.xlsx")

    def escaped(value):
        return str(value).replace("|", "\\|").replace("\n", " ")

    lines = [
        "# Migration QA Evidence",
        "",
        f"Validation: **{verdict}** | Coverage: **{summary['coverage']}**",
        "",
        "Production sign-off: **not granted**. Evidence is scoped to the named run and selected stages.",
        "",
        "| Mapping | Stage | Status | Source | Target | Mismatches | Seconds |",
        "|---|---|---|---:|---:|---:|---:|",
    ]
    for row in results:
        lines.append(
            "| "
            + " | ".join(
                escaped(row.get(k, "-"))
                for k in (
                    "mapping",
                    "stage",
                    "status",
                    "source_count",
                    "target_count",
                    "mismatches",
                    "seconds",
                )
            )
            + " |"
        )
    lines += [
        "",
        "## Cutover",
        "",
        f"- snapshot_id: {cutover['snapshot_id'] or '-'}",
        f"- enforced: {cutover['snapshot_enforced']} ({cutover['source_snapshot'] or 'recorded only'})",
    ]
    if cutover["post_snapshot_arrivals"]:
        lines.append(f"- post-snapshot arrivals: {cutover['post_snapshot_arrivals']}")
    if cutover["target_asof"]:
        lines.append(f"- target as-of max: {cutover['target_asof']}")
    lines += ["", "## Insights and improvement suggestions", ""] + [
        f"- {item}" for item in insights
    ]
    lines += [
        "",
        "## Evidence",
        "",
        "- `run_summary.json`: complete per-mapping stage results and run metadata.",
        "- `migration_summary.xlsx`: audit workbook, metadata, bounded redacted samples.",
        "- `output.xml`, `report.html`, `log.html`: Robot execution evidence.",
    ]
    (out / "migration_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return summary["verdict"]
