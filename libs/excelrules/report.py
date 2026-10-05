"""Write the Excel summary report: a Run Info audit sheet (timestamp,
environment, contract, commit, verdict), one row per rule with status,
violation count, and a diff sample, plus a full Violations detail sheet —
the audit artifact a reviewer opens next to the Robot HTML reports."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import RuleResult

_FILL = {"PASS": "C6EFCE", "WARN": "FFEB9C", "FAIL": "FFC7CE", "ERROR": "FFC7CE"}


def _bold_header(ws) -> None:
    for c in ws[1]:
        c.font = Font(bold=True)


def write_summary(
    results: list[RuleResult], out_path: "str | Path", meta: "dict | None" = None
) -> Path:
    meta = meta or {}
    wb = Workbook()

    info = wb.active
    info.title = "Run Info"
    info.append(["Field", "Value"])
    _bold_header(info)
    counts = {"PASS": 0, "WARN": 0, "FAIL": 0, "ERROR": 0}
    for r in results:
        counts[r.status] = counts.get(r.status, 0) + 1
    overall = "FAIL" if counts["FAIL"] or counts["ERROR"] else "PASS"
    info_rows = [
        ("Run Timestamp", meta.get("timestamp", "")),
        ("Environment", meta.get("environment", "")),
        ("Contract", meta.get("contract", "")),
        ("Contract Version", meta.get("contract_version", "")),
        ("Rules Workbook", meta.get("rules_workbook", "")),
        ("Executed By", meta.get("executed_by", "")),
        ("Git Commit", meta.get("git_commit", "")),
        ("Rules Executed", len(results)),
        ("Pass", counts["PASS"]),
        ("Warn", counts["WARN"]),
        ("Fail", counts["FAIL"]),
        ("Error", counts["ERROR"]),
        ("Total Violations", sum(len(r.violations) for r in results)),
        ("Overall Status", overall),
    ]
    for field, value in info_rows:
        info.append([field, value])
        info.cell(row=info.max_row, column=1).font = Font(bold=True)
    info.cell(row=info.max_row, column=2).fill = PatternFill("solid", fgColor=_FILL[overall])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 60

    ws = wb.create_sheet("Summary")
    ws.append(
        ["Test ID", "Rule Type", "Target", "Severity", "Status", "Violations", "Sample", "Error"]
    )
    _bold_header(ws)

    for r in results:
        rule = r.rule
        sample = "; ".join(
            f"{v.column} key={v.key}: expected {v.expected!r} actual {v.actual!r}"
            for v in r.violations[:5]
        )
        ws.append(
            [
                rule.test_id,
                rule.rule_type,
                f"{rule.target_table}.{rule.target_column}",
                rule.severity,
                r.status,
                len(r.violations),
                sample,
                r.error or "",
            ]
        )
        status_cell = ws.cell(row=ws.max_row, column=5)
        status_cell.fill = PatternFill("solid", fgColor=_FILL.get(r.status, "FFFFFF"))

    for i, width in enumerate([10, 16, 34, 10, 8, 11, 80, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    det = wb.create_sheet("Violations")
    det.append(["Test ID", "Severity", "Status", "Column", "Key", "Expected", "Actual"])
    _bold_header(det)
    for r in results:
        for v in r.violations:
            det.append(
                [
                    r.rule.test_id,
                    r.rule.severity,
                    r.status,
                    v.column,
                    str(v.key),
                    str(v.expected),
                    str(v.actual),
                ]
            )
    for i, width in enumerate([10, 10, 8, 24, 20, 40, 40], start=1):
        det.column_dimensions[get_column_letter(i)].width = width
    det.freeze_panes = "A2"

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
