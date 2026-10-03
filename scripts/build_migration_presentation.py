"""Build the management/client presentation deck for migration verification.

    python scripts/build_migration_presentation.py

Outputs docs/migration/Webster_Santander_Migration_Verification.pptx
Requires python-pptx (dev dependency — not needed by the framework itself).
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = (
    Path(__file__).resolve().parents[1]
    / "docs"
    / "migration"
    / "Webster_Santander_Migration_Verification.pptx"
)

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
LIGHT = RGBColor(0xDE, 0xEA, 0xF6)
GREEN = RGBColor(0x54, 0x82, 0x35)
GREY = RGBColor(0x40, 0x40, 0x40)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xBF, 0x60, 0x00)

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

# shared defaults (Length objects are immutable ints)
_L, _T, _W, _H, _SZ = Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.4), 16


def _text(shape, lines, size=16, color=GREY, bold_first=False, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(line, tuple):
            text, indent = line
        else:
            text, indent = line, 0
        p.text = text
        p.level = indent
        p.alignment = align
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.color.rgb = color
            run.font.name = "Calibri"
            run.font.bold = bold_first and i == 0


def slide(prs, title, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    _text(bar, [title], size=28, color=WHITE, bold_first=True)
    bar.text_frame.margin_left = Inches(0.5)
    bar.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    if subtitle:
        box = s.shapes.add_textbox(Inches(0.5), Inches(1.12), SLIDE_W - Inches(1), Inches(0.4))
        _text(box, [subtitle], size=14, color=BLUE)
    return s


def bullets(s, lines, left=_L, top=_T, width=_W, height=_H, size=_SZ):
    box = s.shapes.add_textbox(left, top, width, height)
    _text(box, lines, size=size)
    for p in box.text_frame.paragraphs:
        p.space_after = Pt(8)
    return box


def card(s, left, top, width, height, title, lines, fill=LIGHT, title_color=NAVY, size=13):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = BLUE
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = title
    p.runs[0].font.bold = True
    p.runs[0].font.size = Pt(size + 2)
    p.runs[0].font.color.rgb = title_color
    for line in lines:
        para = tf.add_paragraph()
        para.text = line
        para.runs[0].font.size = Pt(size)
        para.runs[0].font.color.rgb = GREY
    return shape


def table_slide(s, headers, rows, left, top, width, height, font=12, widths=None):
    shape = s.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height)
    t = shape.table
    if widths:
        for i, w in enumerate(widths):
            t.columns[i].width = Emu(w)
    for j, h in enumerate(headers):
        cell = t.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for r in cell.text_frame.paragraphs[0].runs:
            r.font.bold = True
            r.font.size = Pt(font)
            r.font.color.rgb = WHITE
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            cell = t.cell(i, j)
            cell.text = str(v)
            for r in cell.text_frame.paragraphs[0].runs:
                r.font.size = Pt(font)
                r.font.color.rgb = GREY
    return t


def flow_arrow(s, x1, y1, x2, y2):
    """Connector with a triangular arrowhead."""
    conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = BLUE
    conn.line.width = Pt(2.25)
    from lxml import etree
    from pptx.oxml.ns import qn

    ln = conn.line._get_or_add_ln()
    ln.append(etree.SubElement(ln, qn("a:tailEnd")))
    ln[-1].set("type", "triangle")
    return conn


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    # ---- 1. Title ---------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    box = s.shapes.add_textbox(Inches(1), Inches(2.3), Inches(11.3), Inches(3))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (txt, sz, col) in enumerate(
        [
            ("Webster → Santander Data Migration", 40, WHITE),
            ("Automated Verification Framework", 30, LIGHT),
            (
                "Contract-driven · Read-only · Auditable — one PASS/FAIL per migrated table",
                16,
                LIGHT,
            ),
        ]
    ):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt
        p.runs[0].font.size = Pt(sz)
        p.runs[0].font.color.rgb = col
        p.runs[0].font.bold = i == 0
    foot = s.shapes.add_textbox(Inches(1), Inches(6.4), Inches(11.3), Inches(0.6))
    _text(
        foot,
        ["QA Automation — reconciliation testing · Management & client review"],
        size=12,
        color=LIGHT,
    )

    # ---- 2. Challenge ------------------------------------------------------
    s = slide(prs, "The problem we are solving", "Why a migration cannot be signed off manually")
    bullets(
        s,
        [
            (
                "Webster → Santander migrates dozens of tables (customers, accounts, transactions, …)",
                0,
            ),
            ("Manual spreadsheet comparison does not scale, cannot be audited, and exposes PII", 0),
            (
                "A single missed row, wrong code mapping, or orphan reference is a regulatory issue",
                0,
            ),
            ("Each party needs independent proof — not just “the ETL ran”", 0),
            "",
            (
                "Our answer: the migration ETL loads the data — this framework independently verifies it",
                0,
            ),
            (
                "with read-only access on both sides and produces a signed, tamper-evident report.",
                0,
            ),
        ],
        size=18,
    )

    # ---- 3. Architecture ---------------------------------------------------
    s = slide(
        prs,
        "How it works",
        "Verification layer — read-only on both sides, separate from the load pipeline",
    )
    card(
        s,
        Inches(0.5),
        Inches(2.0),
        Inches(3.4),
        Inches(1.6),
        "Webster — S3 extracts",
        ["Versioned CSV part-files", "per table prefix", "read via GetObject only"],
    )
    card(
        s,
        Inches(4.95),
        Inches(2.0),
        Inches(3.4),
        Inches(1.6),
        "Bank migration ETL",
        ["Owned by migration team", "Loads + transforms", "(not part of this framework)"],
        fill=WHITE,
    )
    card(
        s,
        Inches(9.4),
        Inches(2.0),
        Inches(3.4),
        Inches(1.6),
        "Santander — Snowflake",
        ["Target tables", "SELECT-only verification role", "key-pair auth, no secondary roles"],
    )
    flow_arrow(s, Inches(3.9), Inches(2.8), Inches(4.95), Inches(2.8))
    flow_arrow(s, Inches(8.35), Inches(2.8), Inches(9.4), Inches(2.8))
    card(
        s,
        Inches(3.2),
        Inches(4.6),
        Inches(6.9),
        Inches(1.7),
        "Verification framework (this project)",
        [
            "Read-only access to BOTH sides — never writes or loads",
            "Manifest of tables → contracts define columns, keys, rules, mappings, controls",
            "One Robot Framework test per table per check level → audited evidence",
        ],
        fill=LIGHT,
    )
    flow_arrow(s, Inches(4.4), Inches(4.6), Inches(3.2), Inches(3.6))
    flow_arrow(s, Inches(8.5), Inches(4.6), Inches(10.0), Inches(3.6))

    # ---- 4. Contract-driven config ------------------------------------------
    s = slide(
        prs,
        "Config-driven, not code-driven",
        "Adding table 51 = one contract + one manifest line — zero suite changes",
    )
    card(
        s,
        Inches(0.5),
        Inches(1.7),
        Inches(6.0),
        Inches(2.4),
        "Migration manifest",
        [
            "Tables in scope, criticality tier, dependencies",
            "Cross-table relationships to verify",
            "Data classification gates which environments may run",
        ],
    )
    card(
        s,
        Inches(6.8),
        Inches(1.7),
        Inches(6.0),
        Inches(2.4),
        "Per-table contract (YAML)",
        [
            "Column mapping Webster → Santander, data types, keys",
            "DQ rules, code mappings (e.g. status A→ACTIVE), control totals",
            "pii: true flags drive masking",
        ],
    )
    bullets(
        s,
        [
            ("Robot pre-run expansion generates the test set from the manifest —", 0),
            ("3 tables + 2 relationships already produce 23 distinct checks today", 1),
            ("Same contracts run in every environment (mock / SIT / UAT / production)", 0),
        ],
        top=Inches(4.5),
        size=16,
    )

    # ---- 5. Verification levels ----------------------------------------------
    s = slide(
        prs, "Six levels of proof per table", "Cheapest checks first; record-level proof on top"
    )
    table_slide(
        s,
        ["Level", "Check", "How"],
        [
            ["L0", "Source lineage — exact S3 objects read (key, version, ETag)", "S3 API"],
            ["L1", "Source data quality + code-mapping completeness", "Extract, in-memory"],
            [
                "L2",
                "Target schema — columns, types, nullability, primary key",
                "INFORMATION_SCHEMA",
            ],
            [
                "L3",
                "Row counts + control totals (sums, min/max, group-by)",
                "Aggregated inside Snowflake",
            ],
            [
                "L4",
                "Referential integrity — no orphan accounts/transactions",
                "Anti-join inside Snowflake",
            ],
            [
                "L5",
                "Record-level comparison incl. transformations",
                "Bounded, fail-closed above limit",
            ],
        ],
        Inches(0.5),
        Inches(1.7),
        Inches(12.3),
        Inches(4.6),
        font=14,
        widths=[Inches(1.2), Inches(7.4), Inches(3.7)],
    )

    # ---- 6. Security ----------------------------------------------------------
    s = slide(prs, "Security controls built in", "Designed for a regulated banking environment")
    cards = [
        (
            "Read-only by design",
            [
                "Snowflake: dedicated SELECT-only role, secondary roles disabled",
                "S3: GetObject / ListBucket only",
                "Framework cannot write — by code and by RBAC",
            ],
        ),
        (
            "No PII in evidence",
            [
                "pii:true values → ***REDACTED***, or keyed HMAC tokens",
                "So failures correlate across tables without exposing data",
                "Verified: zero raw PII in reports/logs",
            ],
        ),
        (
            "Fail closed",
            [
                "Missing credentials, empty extracts, unknown codes,",
                "drifted headers, row-limit breach, wrong role → FAIL",
                "A partial run can only be INCOMPLETE — never PASS",
            ],
        ),
        (
            "Tamper-evident audit",
            [
                "Every run: id, operator, git commit, SHA-256 of rules",
                "and S3 object versions; report carries its own hash",
                "Store in immutable evidence store for sign-off",
            ],
        ),
    ]
    for i, (title, lines) in enumerate(cards):
        card(
            s,
            Inches(0.5 + (i % 2) * 6.4),
            Inches(1.7 + (i // 2) * 2.5),
            Inches(6.0),
            Inches(2.25),
            title,
            lines,
        )

    # ---- 7. Defects proven to be caught ----------------------------------------
    s = slide(
        prs,
        "Proven to catch the defects that matter",
        "Negative-path suite injects real faults — every one is detected",
    )
    table_slide(
        s,
        ["Injected defect", "Detected by"],
        [
            [
                "Account balance changed by +0.01",
                "Control totals (balance_by_currency) + record compare",
            ],
            ["Target row deleted", "Row count + missing-key + control count"],
            ["Orphan account (no customer)", "Referential integrity — masked in output"],
            ["Unmapped source code in extract", "Data-quality mapping rule"],
            ["S3 part file with wrong header", "Lineage/header check — run fails closed"],
            ["PII in failure output", "Never — assertion scans prove masking"],
        ],
        Inches(0.5),
        Inches(1.7),
        Inches(12.3),
        Inches(4.8),
        font=14,
        widths=[Inches(5.2), Inches(7.1)],
    )

    # ---- 8. Results ------------------------------------------------------------
    s = slide(
        prs,
        "Current verification status",
        "Executed on synthetic bank-like data — Webster extract stub + targets",
    )
    table_slide(
        s,
        ["Evidence", "Result"],
        [
            ["Migration suite — 3 tables, 2 relationships, 23 generated checks", "23 / 23 PASS"],
            ["Negative-path proofs (defect injection)", "7 / 7 PASS"],
            ["Regression — CSV/API/S3/MySQL/Postgres suites", "63 / 63 PASS"],
            ["Unit tests (masking, controls, manifest, S3, audit)", "93 / 93 PASS"],
            ["Live Snowflake trial (key-pair, read-only role)", "14 / 14 PASS"],
            ["Lint / types / secrets scanning (ruff, mypy, gitleaks)", "Clean"],
        ],
        Inches(0.5),
        Inches(1.7),
        Inches(12.3),
        Inches(4.6),
        font=14,
        widths=[Inches(8.6), Inches(3.7)],
    )
    bullets(
        s,
        [("All demo data synthetic — no real customer data has entered this framework.", 0)],
        top=Inches(6.5),
        size=12,
    )

    # ---- 9. Roadmap -------------------------------------------------------------
    s = slide(prs, "Roadmap to production sign-off", "Phase 1 done; Phase 2 needs bank inputs")
    card(
        s,
        Inches(0.5),
        Inches(1.7),
        Inches(6.0),
        Inches(4.6),
        "Phase 1 — delivered",
        [
            "Manifest-driven per-table verification",
            "PII masking, lineage, audit report",
            "Control totals + referential integrity",
            "Code mappings, S3 hardening",
            "CI pipeline + evidence artifacts",
        ],
        fill=LIGHT,
    )
    card(
        s,
        Inches(6.8),
        Inches(1.7),
        Inches(6.0),
        Inches(4.6),
        "Phase 2 — in progress",
        [
            "Contract generation from approved mapping spec",
            "Chunked hash comparison for >100k-row tables",
            "Balance proofs, control/trailer files, delta runs",
            "Parallel execution + central evidence store",
            "Bank-approved run: real S3 → real Snowflake",
        ],
        fill=WHITE,
    )

    # ---- 10. Asks --------------------------------------------------------------
    s = slide(prs, "What we need for go-live", "Approvals and inputs requested from the client")
    bullets(
        s,
        [
            ("Approved source-to-target mapping spec (tables, columns, code lists)", 0),
            ("Read-only access: S3 extract bucket + Snowflake verification role (key-pair)", 0),
            ("Data classification sign-off — which columns are PII/confidential", 0),
            ("Required control totals per table (e.g. balance by currency/product)", 0),
            ("A bank-approved runner host + evidence retention policy", 0),
            ("Mock-migration schedule — we verify each load until cutover sign-off", 0),
        ],
        size=18,
    )
    card(
        s,
        Inches(0.5),
        Inches(6.0),
        Inches(12.3),
        Inches(1.0),
        "Decision requested",
        ["Approve the verification approach + provide the inputs above to start Phase 2."],
        fill=LIGHT,
        title_color=GREEN,
    )

    # ---- 11. Close ----------------------------------------------------------------
    s = slide(prs, "Summary")
    bullets(
        s,
        [
            ("Verification is independent of the load: if the ETL says it loaded, we prove it.", 0),
            (
                "Six check levels per table, config-driven, scales by adding contracts — not code.",
                0,
            ),
            ("PII never leaves the data plane; every report is hashed and auditable.", 0),
            (
                "Every component is already green on synthetic fixtures and a live Snowflake trial.",
                0,
            ),
            ("Next: Phase 2 inputs from the client, then mock-migration verification runs.", 0),
        ],
        size=18,
    )

    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
