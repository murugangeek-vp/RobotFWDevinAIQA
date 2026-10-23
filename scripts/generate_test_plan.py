"""Generate the Phase-1 test plan workbook: docs/test_plan.xlsx.

Deterministic — derived from the mapping workbooks themselves so the plan
always matches what the suites actually execute. One row per requirement
x check: Requirement ID, dimension (DC/DA/DI/DG), scenario, test case and
the lifecycle test that executes it.

Run: .venv/Scripts/python.exe scripts/generate_test_plan.py
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from libs.mapping.parser import load_workbook  # noqa: E402

WORKBOOKS = [
    ("hop1", REPO / "data" / "mapping" / "phase1_hop1_csv_to_stage.xlsx"),
    ("hop2", REPO / "data" / "mapping" / "phase1_hop2_stage_to_il.xlsx"),
    ("hop3", REPO / "data" / "mapping" / "phase1_hop3_il_to_consumption.xlsx"),
]

_LIFECYCLE = {
    "mapping": "Requirements Intake And Mapping Compile",
    "source": "Source And Schema Conformance",
    "dq": "Source And Schema Conformance",
    "schema": "Source And Schema Conformance",
    "recon": "Transformation And Record Reconciliation",
    "transform": "Transformation And Record Reconciliation",
    "controls": "Business Validation Controls Relations Profile And Rules",
    "relations": "Business Validation Controls Relations Profile And Rules",
    "profile": "Business Validation Controls Relations Profile And Rules",
    "rules": "Business Validation Controls Relations Profile And Rules",
}


def _dim(req: "str | None") -> str:
    for part in (req or "").split("-"):
        if part in ("DC", "DA", "DI", "DG"):
            return part
    return ""


def _rows():
    for hop, path in WORKBOOKS:
        tables, _ = load_workbook(path)
        for t in tables:
            if not t.enabled:
                continue
            req = t.requirement_id or f"REQ-{hop.upper()}-{t.mapping_id}"
            base = [req, _dim(t.requirement_id), hop]
            tgt = f"{t.target_schema}.{t.target_table}"
            yield base + [
                f"{tgt} loads and validates end-to-end",
                "mapping compile",
                _LIFECYCLE["mapping"],
            ]
            yield base + [
                f"{tgt} schema, nullability and source DQ conform",
                "source + schema conformance",
                _LIFECYCLE["schema"],
            ]
            yield base + [
                f"{tgt} holds every source row exactly once",
                "record reconciliation (missing/extra/diffs)",
                _LIFECYCLE["recon"],
            ]
            transforms = [f for f in t.fields if f.enabled and f.transformation]
            if transforms:
                yield base + [
                    f"{len(transforms)} transformation(s) recompute on target",
                    "transformation re-computation",
                    _LIFECYCLE["transform"],
                ]
            for c in t.controls:
                yield [
                    c.get("requirement_id") or req,
                    _dim(c.get("requirement_id")),
                    hop,
                    f"control {c['id']} ({c['type']}) certified vs trailer",
                    "trailer-certified aggregate control",
                    _LIFECYCLE["controls"],
                ]
            for rel in t.relations:
                yield [
                    rel.get("requirement_id") or req,
                    _dim(rel.get("requirement_id")),
                    hop,
                    f"{rel['child_column']} references {rel['parent_mapping']}.{rel['parent_column']}",
                    "FK relation — no orphans",
                    _LIFECYCLE["relations"],
                ]
        try:
            from libs.excelrules import parser as rules_parser

            for r in rules_parser.load_rules(path):
                if not r.enabled:
                    continue
                yield [
                    r.requirement_id or f"REQ-{hop.upper()}-RULES",
                    _dim(r.requirement_id),
                    hop,
                    f"{r.test_id}: {r.business_rule or r.rule_type}",
                    f"business rule ({r.rule_type})",
                    _LIFECYCLE["rules"],
                ]
        except ValueError:
            continue  # workbook without a Rules sheet


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Plan"
    ws.append(["Requirement ID", "Dimension", "Hop", "Requirement", "Scenario", "Test Case"])
    for c in ws[1]:
        c.font = Font(bold=True)
    n = 0
    for row in _rows():
        ws.append(row)
        n += 1
    for col, w in zip("ABCDEF", [16, 10, 6, 60, 44, 52]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    out = REPO / "docs" / "test_plan.xlsx"
    wb.save(out)
    print(f"wrote {out} ({n} test-case rows)")


if __name__ == "__main__":
    main()
