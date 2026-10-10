"""Build data/mapping/bank_mapping.xlsx — the Webster->Santander table and
field mapping workbook. This IS the template QA fills in per migration:
one Tables row per source->target table pair, one FieldMap row per mapped
field (rename or engine transform expression), plus the Mapping sheet of
code-translation tables. Deterministic output so the artifact is auditable;
the committed .xlsx is what the suites read.

Run: .venv/Scripts/python.exe scripts/generate_mapping_workbook.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

REPO = Path(__file__).resolve().parent.parent
OUT = REPO / "data" / "mapping" / "bank_mapping.xlsx"

TABLES_HEADERS = [
    "Mapping ID",
    "Source Schema",
    "Source Table",
    "Target Schema",
    "Target Table",
    "Key Columns",
    "Where Clause",
    "Recon Mode",
    "Batch Size",
    "Enabled",
]

# key columns are TARGET-side names (they must have enabled FieldMap rows).
TABLES = [
    [
        "customer_migration",
        "webster",
        "customer",
        "PUBLIC",
        "customer",
        "customer_id",
        "",
        "full",
        "",
        "yes",
    ],
    [
        "account_migration",
        "webster",
        "account",
        "PUBLIC",
        "account",
        "account_id",
        "",
        "full",
        "",
        "yes",
    ],
    [
        "transaction_migration",
        "webster",
        "txn",
        "PUBLIC",
        "account_txn",
        "txn_id",
        "",
        "full",
        "",
        "yes",
    ],
    ["loan_migration", "webster", "loan", "PUBLIC", "loan", "loan_id", "", "full", "", "yes"],
    [
        "branch_migration",
        "webster",
        "branch",
        "PUBLIC",
        "branch",
        "branch_id",
        "",
        "full",
        "",
        "yes",
    ],
    # Disabled example — proves Enabled=no rows are parsed but skipped:
    ["legacy_card_migration", "webster", "card", "PUBLIC", "card", "card_id", "", "full", "", "no"],
]

FIELDMAP_HEADERS = [
    "Mapping ID",
    "Source Column",
    "Target Column",
    "Target Type",
    "Nullable",
    "Max Length",
    "Scale",
    "Transformation",
    "PII",
    "Enabled",
]

# PII=yes columns get their diff sample values redacted in evidence;
# non-PII columns show bounded actual values for reviewers.
_PII_TARGETS = {"ssn_last4", "phone", "email", "full_name", "address"}

# Transformation: blank = direct move (rename when names differ); otherwise an
# engine expression — {col} placeholders, upper/lower/trim/substr/round/
# date_fmt/ifnull/casewhen, or map(src_col, mapping_name) against the Mapping
# sheet. Source Column is the declared primary input (also drives the source
# SELECT); extra columns referenced by {placeholders} are fetched too.
FIELDMAP = [
    # ---- customer_migration: webster.customer -> PUBLIC.customer
    ["customer_migration", "cust_no", "customer_id", "integer", "no", "", "", "", "yes"],
    ["customer_migration", "fname", "first_name", "string", "no", "60", "", "", "yes"],
    ["customer_migration", "lname", "last_name", "string", "no", "60", "", "", "yes"],
    [
        "customer_migration",
        "fname",
        "full_name",
        "string",
        "no",
        "201",
        "",
        "{fname} {lname}",
        "yes",
    ],
    ["customer_migration", "email", "email", "string", "no", "320", "", "lower({email})", "yes"],
    [
        "customer_migration",
        "dob_txt",
        "dob",
        "date",
        "no",
        "",
        "",
        "date_fmt({dob_txt}, '%Y-%m-%d', '%m/%d/%Y')",
        "yes",
    ],
    [
        "customer_migration",
        "ssn",
        "ssn_last4",
        "string",
        "yes",
        "4",
        "",
        "substr({ssn}, 7, 4)",
        "yes",
    ],
    ["customer_migration", "addr1", "address", "string", "no", "120", "", "", "yes"],
    ["customer_migration", "city", "city", "string", "no", "60", "", "upper({city})", "yes"],
    ["customer_migration", "st", "state", "string", "no", "2", "", "", "yes"],
    ["customer_migration", "zip", "zip_code", "string", "no", "10", "", "", "yes"],
    ["customer_migration", "phone", "phone", "string", "yes", "25", "", "", "yes"],
    [
        "customer_migration",
        "stat_cd",
        "status",
        "string",
        "no",
        "10",
        "",
        "map(stat_cd, status_codes)",
        "yes",
    ],
    ["customer_migration", "open_dt", "open_date", "date", "no", "", "", "", "yes"],
    # ---- account_migration: webster.account -> PUBLIC.account
    ["account_migration", "acct_no", "account_id", "integer", "no", "", "", "", "yes"],
    ["account_migration", "cust_no", "customer_id", "integer", "no", "", "", "", "yes"],
    ["account_migration", "prod_cd", "product_code", "string", "no", "3", "", "", "yes"],
    [
        "account_migration",
        "prod_cd",
        "product_name",
        "string",
        "no",
        "30",
        "",
        "map(prod_cd, product_codes)",
        "yes",
    ],
    ["account_migration", "branch_no", "branch_id", "integer", "no", "", "", "", "yes"],
    ["account_migration", "open_dt", "open_date", "date", "no", "", "", "", "yes"],
    [
        "account_migration",
        "bal_amt",
        "balance",
        "decimal",
        "no",
        "",
        "2",
        "round({bal_amt}, 2)",
        "yes",
    ],
    [
        "account_migration",
        "stat_cd",
        "status",
        "string",
        "no",
        "10",
        "",
        "map(stat_cd, status_codes)",
        "yes",
    ],
    # ---- transaction_migration: webster.txn -> PUBLIC.account_txn
    ["transaction_migration", "txn_id", "txn_id", "integer", "no", "", "", "", "yes"],
    ["transaction_migration", "acct_no", "account_id", "integer", "no", "", "", "", "yes"],
    ["transaction_migration", "txn_dt", "txn_date", "date", "no", "", "", "", "yes"],
    ["transaction_migration", "txn_tm", "txn_time", "string", "yes", "8", "", "", "yes"],
    [
        "transaction_migration",
        "txn_cd",
        "txn_type",
        "string",
        "no",
        "20",
        "",
        "map(txn_cd, txn_codes)",
        "yes",
    ],
    [
        "transaction_migration",
        "amount",
        "amount",
        "decimal",
        "no",
        "",
        "2",
        "round({amount}, 2)",
        "yes",
    ],
    [
        "transaction_migration",
        "desc_txt",
        "description",
        "string",
        "yes",
        "100",
        "",
        "trim({desc_txt})",
        "yes",
    ],
    # ---- loan_migration: webster.loan -> PUBLIC.loan
    ["loan_migration", "loan_no", "loan_id", "integer", "no", "", "", "", "yes"],
    ["loan_migration", "cust_no", "customer_id", "integer", "no", "", "", "", "yes"],
    [
        "loan_migration",
        "loan_type",
        "loan_type",
        "string",
        "no",
        "20",
        "",
        "map(loan_type, loan_types)",
        "yes",
    ],
    [
        "loan_migration",
        "orig_amt",
        "original_amount",
        "decimal",
        "no",
        "",
        "2",
        "round({orig_amt}, 2)",
        "yes",
    ],
    [
        "loan_migration",
        "curr_bal",
        "current_balance",
        "decimal",
        "no",
        "",
        "2",
        "round({curr_bal}, 2)",
        "yes",
    ],
    ["loan_migration", "rate_pct", "rate", "decimal", "no", "", "3", "round({rate_pct}, 3)", "yes"],
    ["loan_migration", "start_dt", "start_date", "date", "no", "", "", "", "yes"],
    ["loan_migration", "mat_dt", "maturity_date", "date", "yes", "", "", "", "yes"],
    [
        "loan_migration",
        "stat_cd",
        "status",
        "string",
        "no",
        "10",
        "",
        "map(stat_cd, status_codes)",
        "yes",
    ],
    # ---- branch_migration: webster.branch -> PUBLIC.branch
    ["branch_migration", "branch_no", "branch_id", "integer", "no", "", "", "", "yes"],
    [
        "branch_migration",
        "branch_nm",
        "branch_name",
        "string",
        "no",
        "80",
        "",
        "trim({branch_nm})",
        "yes",
    ],
    ["branch_migration", "city", "city", "string", "no", "60", "", "upper({city})", "yes"],
    ["branch_migration", "st", "state", "string", "no", "2", "", "", "yes"],
    [
        "branch_migration",
        "region_cd",
        "region",
        "string",
        "no",
        "20",
        "",
        "map(region_cd, region_codes)",
        "yes",
    ],
]

CONTROLS_HEADERS = [
    "Mapping ID",
    "Control ID",
    "Type",
    "Column",
    "Group By",
    "File Name",
    "Enabled",
]

# Aggregate controls — the financial-totals layer. Type: count | sum | min |
# max | null_count | count_distinct. Column/Group By name TARGET-side
# contract columns. File Name is the certified total's name in the bank
# trailer file (source.control_file); absent = the control id is used.
CONTROLS = [
    ["customer_migration", "customer_rows", "count", "", "", "customer_rows", "yes"],
    ["customer_migration", "null_phone", "null_count", "phone", "", "null_phone", "yes"],
    ["account_migration", "account_rows", "count", "", "", "account_rows", "yes"],
    [
        "account_migration",
        "balance_total",
        "sum",
        "balance",
        "",
        "balance_total",
        "yes",
    ],
    [
        "account_migration",
        "balance_by_branch",
        "sum",
        "balance",
        "branch_id",
        "balance_by_branch",
        "yes",
    ],
    [
        "account_migration",
        "distinct_customers",
        "count_distinct",
        "customer_id",
        "",
        "distinct_customers",
        "yes",
    ],
    ["transaction_migration", "txn_rows", "count", "", "", "txn_rows", "yes"],
    [
        "transaction_migration",
        "txn_amount_total",
        "sum",
        "amount",
        "",
        "txn_amount_total",
        "yes",
    ],
    [
        "transaction_migration",
        "txn_by_type",
        "count",
        "",
        "txn_type",
        "txn_by_type",
        "yes",
    ],
    ["transaction_migration", "max_txn", "max", "amount", "", "max_txn", "yes"],
    ["loan_migration", "loan_rows", "count", "", "", "loan_rows", "yes"],
    [
        "loan_migration",
        "loan_balance_total",
        "sum",
        "current_balance",
        "",
        "loan_balance_total",
        "yes",
    ],
    ["branch_migration", "branch_rows", "count", "", "", "branch_rows", "yes"],
]

RELATIONS_HEADERS = [
    "Parent Mapping",
    "Parent Column",
    "Child Mapping",
    "Child Column",
    "Enabled",
]

# FK integrity across mappings — target-side column names on both ends.
# The relations stage checks both sides: orphans in source space (source DQ
# problem) and in the target (a migration integrity break). Null FKs are
# counted separately, not treated as orphans.
RELATIONS = [
    ["customer_migration", "customer_id", "account_migration", "customer_id", "yes"],
    ["branch_migration", "branch_id", "account_migration", "branch_id", "yes"],
    ["account_migration", "account_id", "transaction_migration", "account_id", "yes"],
    ["customer_migration", "customer_id", "loan_migration", "customer_id", "yes"],
]

MAPPING_HEADERS = ["Mapping Name", "Code", "Value"]
MAPPING = [
    ["status_codes", "A", "active"],
    ["status_codes", "I", "inactive"],
    ["status_codes", "S", "suspended"],
    ["product_codes", "CHK", "Checking"],
    ["product_codes", "SAV", "Savings"],
    ["product_codes", "MMA", "Money Market"],
    ["product_codes", "CD", "Certificate of Deposit"],
    ["txn_codes", "DBT", "debit"],
    ["txn_codes", "CRD", "credit"],
    ["txn_codes", "FEE", "fee"],
    ["txn_codes", "INT", "interest"],
    ["loan_types", "AUTO", "Auto Loan"],
    ["loan_types", "HOME", "Home Loan"],
    ["loan_types", "PERS", "Personal Loan"],
    ["loan_types", "COMM", "Commercial Loan"],
    ["region_codes", "NE", "Northeast"],
    ["region_codes", "SE", "Southeast"],
    ["region_codes", "MW", "Midwest"],
    ["region_codes", "WE", "West"],
]

README = [
    ["Webster -> Santander migration mapping workbook"],
    [""],
    ["Tables sheet — one row per source->target table pair."],
    ["  Mapping ID: unique label for the pair (used to select a run)."],
    ["  Source/Target Schema+Table: physical table names."],
    ["  Key Columns: comma-separated TARGET-side key columns for reconciliation."],
    ["  Where Clause: optional source filter (read-only fragment, e.g. stat_cd = 'A')."],
    ["  Recon Mode: full | hash (hash = key->rowhash fast path for big tables)."],
    ["  Batch Size: rows per fetch chunk for streaming compares (default 50000)."],
    ["  Enabled: no = parsed but skipped."],
    [""],
    ["FieldMap sheet — one row per target field."],
    ["  Source Column: the column fetched from the source table."],
    ["  Transformation: blank = direct move (rename when names differ);"],
    ["    otherwise an engine expression — {col} templates, upper/lower/trim/"],
    ["    substr/round/date_fmt/ifnull/casewhen, or map(src_col, mapping_name)."],
    ["  Target Type: integer | decimal | string | date | timestamp | boolean."],
    ["  Nullable / Max Length / Scale: target-side contract constraints."],
    [""],
    ["Mapping sheet — code translation tables for map() transforms."],
    ["  Mapping Name: referenced by map(col, name); Code -> Value pairs."],
    [""],
    ["Controls sheet — aggregate controls (optional, financial totals)."],
    ["  Type: count | sum | min | max | null_count | count_distinct."],
    ["  Column / Group By name TARGET-side columns (comma-sep, max 3)."],
    ["  File Name: the certified total's name in the bank trailer file"],
    ["    (env source.control_file); blank = Control ID is used."],
    [""],
    ["Relations sheet — FK pairs across mappings (optional)."],
    ["  Parent/Child columns are target-side names; every child value must"],
    ["  exist in the parent's column on BOTH source and target sides."],
    [""],
    ["Connection details are NOT in this workbook — they live in"],
    ["config/environments/*.yaml; credentials come from env vars only."],
]


def _sheet(wb, title, headers, rows, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    if first:
        ws.title = title
    if headers:
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    return ws


def main():
    wb = Workbook()
    _sheet(wb, "README", None, README, first=True)
    wb["README"].column_dimensions["A"].width = 90
    _sheet(wb, "Tables", TABLES_HEADERS, TABLES)
    _sheet(
        wb,
        "FieldMap",
        FIELDMAP_HEADERS,
        [[*r[:8], "yes" if r[2] in _PII_TARGETS else "no", r[8]] for r in FIELDMAP],
    )
    _sheet(wb, "Controls", CONTROLS_HEADERS, CONTROLS)
    _sheet(wb, "Relations", RELATIONS_HEADERS, RELATIONS)
    _sheet(wb, "Mapping", MAPPING_HEADERS, MAPPING)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
