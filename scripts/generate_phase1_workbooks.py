"""Build the three Phase-1 multi-hop mapping workbooks from
data/trailer/Phase1.xlsx semantics:

  hop 1  CSV files                 -> wdr_uat_active_stage.signature
  hop 2  signature (views/tables)  -> wdr_uat_integration.wdr_il_core
  hop 3  wdr_il_core view          -> wdr_uat_consumption.wdr_cl_fds_core

Spec gaps resolved explicitly (documented in each workbook README):

- CUSIC -> APPL_KEY + NAICS_CD (row 10): APPL_KEY is independently hardcoded
  to 'CBS' (row 9 is authoritative); the CUSIC row contributes NAICS_CD only.
- "Complex Logic - details shared in next tab" -> PARTY_TYPE: the next tab
  does not exist. PARTY_TYPE is populated with literal('PENDING_SPEC') so the
  column is verifiably present and the gap is visible in the data + evidence
  rather than silently dropped.
- CASE arms without ELSE (BUSINESS_DOMAIN): ELSE 'UNSPECIFIED' declared.
- DIM_REFERENCE lookup (TAX_ID_NUMBER_TYPE_DESC): resolved through
  signature.v_dim_customer_src — a read-only view joining CUP00301 +
  CUP01201 + DIM_REFERENCE. Hop-3 likewise reads wdr_il_core.v_party_snpsht_src
  = DIM_CUSTOMER LEFT JOIN DIM_CUSTOMER_ADDRESS. Reference data is staged from
  CSV like the other extracts (deviation: spec draws it already inside IL).
- LOAD_DATE: carried as a stage column (spec row 4); run.load_date also
  declared in env for audit.

Run: .venv/Scripts/python.exe scripts/generate_phase1_workbooks.py
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from libs.mapping.sql_transform import sql_to_transform

OUT_DIR = REPO / "data" / "mapping"

# Client-specified CUSTOMER_CATEGORY derivation (spec SQL verbatim except the
# inner CASE's implicit ELSE, declared as ELSE NULL — NULL is the faithful
# result for 'C'/'B' names with no LLC/INC/BANK marker). Translated through
# sql_to_transform so the workbook stores the engine expression.
CUSTOMER_CATEGORY_SQL = (
    "CASE WHEN ref2.attribute_cd IS NOT NULL THEN ref2.attribute_cd "
    "ELSE (CASE WHEN cust.customer_type_cd = 'P' THEN 'INDIVIDUAL' "
    "WHEN cust.customer_type_cd IN ('C', 'B') "
    "AND (UPPER(cust.full_nm) LIKE '%LLC%' OR UPPER(cust.full_nm) LIKE '%INC%') "
    "THEN 'CORP_SMALL' "
    "WHEN cust.customer_type_cd IN ('C', 'B') AND UPPER(cust.full_nm) LIKE '%BANK%' "
    "THEN 'BANK' ELSE NULL END) END"
)
CUSTOMER_CATEGORY_TRANSFORM = sql_to_transform(CUSTOMER_CATEGORY_SQL)

TABLES_HEADERS = [
    "Mapping ID",
    "Source Schema",
    "Source Table",
    "Target Schema",
    "Target Table",
    "Key Columns",
    "Where Clause",
    "Recon Mode",
    "Batch Size",
    "Enabled",
    "Requirement ID",
]
FIELDMAP_HEADERS = [
    "Mapping ID",
    "Source Column",
    "Target Column",
    "Target Type",
    "Nullable",
    "Max Length",
    "Scale",
    "Transformation",
    "PII",
    "Enabled",
    "Business Rule",
    "Requirement ID",
]
MAPPING_HEADERS = ["Mapping Name", "Code", "Value"]
CONTROLS_HEADERS = [
    "Mapping ID",
    "Control ID",
    "Type",
    "Column",
    "Group By",
    "File Name",
    "Enabled",
    "Requirement ID",
]
RELATIONS_HEADERS = [
    "Parent Mapping",
    "Parent Column",
    "Child Mapping",
    "Child Column",
    "Enabled",
    "Requirement ID",
]
RULES_HEADERS = [
    "Test ID",
    "Rule Type",
    "Source Table",
    "Source Column(s)",
    "Target Table",
    "Target Column",
    "Transformation Logic",
    "Expected",
    "Severity",
    "Where Clause",
    "Key Column",
    "Enabled",
    "Business Rule",
    "Requirement ID",
]

# Complex business rules — live on the hop-1 workbook's Rules sheet so the
# end-to-end lifecycle suite needs only the one input template. They exercise the
# excelrules engine's complex-check surface against the hop-1 landed data
# (signature schema) and its read-only join view v_dim_customer_src
# (CUP00301 LEFT JOIN CUP01201 LEFT JOIN DIM_REFERENCE): aggregation via
# inner select, correlated EXISTS, multi-file joins, lookup + default
# logic, conditional and cross-file business rules.
HOP1_RULES = [
    [
        "BR-AGG-UNIQ-CUNBR",
        "CUSTOM_SQL",
        "",
        "",
        "cup00301",
        "cunbr",
        "cunbr IN (SELECT cunbr FROM signature.cup00301 " "GROUP BY cunbr HAVING COUNT(*) = 1)",
        "",
        "CRITICAL",
        "",
        "cunbr",
        "yes",
        "customer number must be unique across the landed customer table",
        "REQ-DI-103",
    ],
    [
        "BR-AGG-BANK-CAP",
        "CUSTOM_SQL",
        "",
        "",
        "cup00301",
        "cubk",
        "(SELECT COUNT(*) FROM signature.cup00301 x WHERE x.cubk = cup00301.cubk) <= 80",
        "",
        "HIGH",
        "",
        "cunbr",
        "yes",
        "no single bank may supply more than 80 landed customers",
        "REQ-DA-104",
    ],
    [
        "BR-FK-CUNBR-EXISTS",
        "CUSTOM_SQL",
        "",
        "",
        "cup01201",
        "cunbr",
        "EXISTS (SELECT 1 FROM signature.cup00301 p WHERE p.cunbr = cup01201.cunbr)",
        "",
        "CRITICAL",
        "",
        "cunbr",
        "yes",
        "every legal-entity row must reference a customer that landed",
        "REQ-DI-102",
    ],
    [
        "BR-TAXID-FORMAT",
        "CUSTOM_SQL",
        "",
        "",
        "cup00301",
        "cussnr",
        "cussty IS NULL OR cussty <> 'S' OR cussnr IS NULL OR cussnr ~ '^[0-9]{9}$'",
        "",
        "HIGH",
        "",
        "cunbr",
        "yes",
        "tax ids declared as SSN-type ('S') must be exactly 9 digits when present",
        "REQ-DI-104",
    ],
    [
        "BR-JOIN-LEGAL-ENTITY",
        "DIRECT_COMPARE",
        "cup00301, cup01201",
        "culegenty",
        "v_dim_customer_src",
        "culegenty",
        "",
        "",
        "HIGH",
        "",
        "cunbr",
        "yes",
        "legal-entity code in the join view must match the CUP01201 extract per customer",
        "REQ-DA-105",
    ],
    [
        "BR-LOOKUP-TAXID-DESC",
        "LOOKUP",
        "cup00301, dim_reference@cussty=attribute_cd",
        "",
        "v_dim_customer_src",
        "attribute_value_desc",
        "lookup(attribute_value_desc)",
        "",
        "HIGH",
        "cussty IS NOT NULL",
        "cunbr",
        "yes",
        "tax-id-type description must match the DIM_REFERENCE decode for declared types",
        "REQ-DA-106",
    ],
    [
        "BR-DEFAULT-TAXID-DESC",
        "CUSTOM_PYTHON",
        "cup00301, cup01201, dim_reference@cussty=attribute_cd",
        "",
        "v_dim_customer_src",
        "attribute_value_desc",
        "attribute_value_desc if attribute_value_desc is not None " "else 'SOCIAL SECURITY NUMBER'",
        "",
        "HIGH",
        "",
        "cunbr",
        "yes",
        "a blank tax-id type decodes to 'SOCIAL SECURITY NUMBER' (spec default)",
        "REQ-DA-106",
    ],
    [
        "BR-EXTRACT-CUT-DATE",
        "CUSTOM_PYTHON",
        "cup00301, cup01201",
        "load_date",
        "cup01201",
        "load_date",
        "load_date",
        "",
        "MEDIUM",
        "",
        "cunbr",
        "yes",
        "legal-entity extract cut date must equal the customer extract cut date",
        "REQ-DC-105",
    ],
    [
        "BR-STATE-FORMAT",
        "REGEX",
        "",
        "",
        "cup00301",
        "cuadst",
        "",
        "^[A-Z]{2}$",
        "MEDIUM",
        "",
        "cunbr",
        "yes",
        "state code must be two uppercase letters",
        "REQ-DI-105",
    ],
    [
        "BR-JOBDATE-FORMAT",
        "DATE_FORMAT",
        "",
        "",
        "cup00301",
        "cujbdt",
        "",
        "%Y%m%d",
        "MEDIUM",
        "",
        "cunbr",
        "yes",
        "birth date must be a valid 8-digit YYYYMMDD value",
        "REQ-DI-106",
    ],
    [
        "BR-SCOPED-NOTNULL",
        "NULL_CHECK",
        "",
        "",
        "cup01201",
        "culegenty",
        "",
        "",
        "HIGH",
        "EXISTS (SELECT 1 FROM signature.cup00301 c WHERE c.cunbr = cup01201.cunbr)",
        "cunbr",
        "yes",
        "legal-entity rows whose customer landed must have a non-null entity type",
        "REQ-DI-107",
    ],
]


def _fm(
    mid,
    src,
    tgt,
    ttype,
    nullable,
    maxlen="",
    scale="",
    transform="",
    pii="no",
    desc="",
    req="",
):
    """One FieldMap row. `desc` is a plain-English explanation of the rule
    (or 'direct copy') for reviewers — documentation only, not executed.
    `req` is the optional Requirement ID used for RTM traceability."""
    return [mid, src, tgt, ttype, nullable, maxlen, scale, transform, pii, "yes", desc, req]


# ---------------------------------------------------------------- hop 1
HOP1_TABLES = [
    [
        "csv_cup00301",
        "csv",
        "cup00301",
        "signature",
        "cup00301",
        "cunbr",
        "",
        "full",
        "",
        "yes",
        "REQ-DC-101",
    ],
    [
        "csv_cup01201",
        "csv",
        "cup01201",
        "signature",
        "cup01201",
        "cunbr",
        "",
        "full",
        "",
        "yes",
        "REQ-DC-102",
    ],
    [
        "csv_dim_reference",
        "csv",
        "dim_reference",
        "signature",
        "dim_reference",
        "attribute_cd",
        "",
        "full",
        "",
        "yes",
        "REQ-DC-103",
    ],
]

HOP1_FIELDMAP = [
    # cup00301 — customer master extract (raw landing: direct copies only)
    _fm("csv_cup00301", "cubk", "cubk", "integer", "no", desc="direct copy"),
    _fm("csv_cup00301", "cunbr", "cunbr", "string", "no", "20", desc="direct copy"),
    _fm("csv_cup00301", "cupers", "cupers", "string", "no", "1", desc="direct copy"),
    _fm("csv_cup00301", "cucncd", "cucncd", "integer", "no", desc="direct copy"),
    _fm("csv_cup00301", "cusic", "cusic", "integer", "yes", desc="direct copy"),
    _fm("csv_cup00301", "cuna1", "cuna1", "string", "no", "40", "", "", "yes", "direct copy"),
    _fm("csv_cup00301", "cucfnm", "cucfnm", "string", "yes", "60", "", "", "yes", "direct copy"),
    _fm("csv_cup00301", "cuclnm", "cuclnm", "string", "yes", "60", "", "", "yes", "direct copy"),
    _fm("csv_cup00301", "cussnr", "cussnr", "string", "yes", "11", "", "", "yes", "direct copy"),
    _fm("csv_cup00301", "cussty", "cussty", "string", "yes", "1", desc="direct copy"),
    _fm("csv_cup00301", "cujbdt", "cujbdt", "string", "yes", "8", desc="direct copy"),
    _fm("csv_cup00301", "cuadl1", "cuadl1", "string", "yes", "80", "", "", "yes", "direct copy"),
    _fm("csv_cup00301", "cuadl2", "cuadl2", "string", "yes", "80", "", "", "yes", "direct copy"),
    _fm("csv_cup00301", "cuadct", "cuadct", "string", "yes", "60", "", "", "yes", "direct copy"),
    _fm("csv_cup00301", "cuadst", "cuadst", "string", "yes", "30", "", "", "yes", "direct copy"),
    _fm("csv_cup00301", "cuadpc", "cuadpc", "string", "yes", "15", "", "", "yes", "direct copy"),
    _fm("csv_cup00301", "load_date", "load_date", "date", "no", desc="direct copy"),
    # cup01201 — legal entity extract
    _fm("csv_cup01201", "cunbr", "cunbr", "string", "no", "20", desc="direct copy"),
    _fm("csv_cup01201", "culegenty", "culegenty", "string", "no", "50", desc="direct copy"),
    _fm("csv_cup01201", "load_date", "load_date", "date", "no", desc="direct copy"),
    # dim_reference — reference codes extract
    _fm(
        "csv_dim_reference",
        "attribute_cd",
        "attribute_cd",
        "string",
        "no",
        "50",
        desc="direct copy",
    ),
    _fm(
        "csv_dim_reference",
        "attribute_value_desc",
        "attribute_value_desc",
        "string",
        "no",
        "256",
        desc="direct copy",
    ),
    _fm("csv_dim_reference", "load_date", "load_date", "date", "no", desc="direct copy"),
    _fm(
        "csv_dim_reference",
        "attribute_map_cd",
        "attribute_map_cd",
        "string",
        "yes",
        "50",
        desc="direct copy",
    ),
]

HOP1_CONTROLS = [
    ["csv_cup00301", "cup00301_rows", "count", "", "", "cup00301_rows", "yes", "REQ-DC-101"],
    ["csv_cup01201", "cup01201_rows", "count", "", "", "cup01201_rows", "yes", "REQ-DC-102"],
    [
        "csv_dim_reference",
        "dim_reference_rows",
        "count",
        "",
        "",
        "dim_reference_rows",
        "yes",
        "REQ-DC-103",
    ],
    [
        "csv_cup00301",
        "cup00301_distinct_cunbr",
        "count_distinct",
        "cunbr",
        "",
        "",
        "yes",
        "REQ-DI-101",
    ],
    ["csv_cup00301", "cup00301_taxid_nulls", "null_count", "cussnr", "", "", "yes", "REQ-DC-104"],
    ["csv_cup00301", "cup00301_naics_sum", "sum", "cusic", "", "", "yes", "REQ-DA-101"],
    ["csv_cup00301", "cup00301_max_country", "max", "cucncd", "", "", "yes", "REQ-DA-102"],
    ["csv_cup00301", "cup00301_by_pers", "count", "", "cupers", "", "yes", "REQ-DA-103"],
]

HOP1_RELATIONS = [
    # every legal-entity row must reference an existing customer
    ["csv_cup00301", "cunbr", "csv_cup01201", "cunbr", "yes", "REQ-DI-102"],
]

HOP1_README = [
    ["Phase-1 hop 1: CSV extracts -> wdr_uat_active_stage.signature"],
    [""],
    ["Validates raw landing: one CSV file per table under env source.dir,"],
    ["exact row counts, schema/types/nullability, FK integrity of cup01201"],
    ["against cup00301, and certified aggregates vs phase1_hop1_trailer.csv:"],
    ["row counts, distinct keys, tax-id nulls, NAICS sum, max country code,"],
    ["and the per-customer-type (CUPERS) grouped count."],
    [""],
    ["dim_reference carries ATTRIBUTE_MAP_CD (lookup key) alongside"],
    ["ATTRIBUTE_CD: tax-id-type rows keep it NULL, NAICS->category rows"],
    ["key on the NAICS code for the hop-3 CUSTOMER_CATEGORY join."],
    [""],
    ["Direct copies only — all transformation happens at hop 2 per spec."],
    [""],
    ["Rules sheet: 11 complex business rules (aggregation via inner select,"],
    ["correlated EXISTS, multi-file joins, lookup + default logic, format"],
    ["checks) run by the end-to-end lifecycle suite — this workbook is the single"],
    ["input template. 'Business Rule' columns explain each rule in plain"],
    ["English (documentation only — not executed)."],
]


# ---------------------------------------------------------------- hop 2
HOP2_TABLES = [
    [
        "dim_customer_migration",
        "signature",
        "v_dim_customer_src",
        "wdr_il_core",
        "dim_customer",
        "involved_party_key",
        "",
        "full",
        "",
        "yes",
        "REQ-DA-201",
    ],
    [
        "dim_customer_address_migration",
        "signature",
        "cup00301",
        "wdr_il_core",
        "dim_customer_address",
        "involved_party_key",
        "",
        "full",
        "",
        "yes",
        "REQ-DA-202",
    ],
    [
        "dim_reference_migration",
        "signature",
        "dim_reference",
        "wdr_il_core",
        "dim_reference",
        "attribute_cd",
        "",
        "full",
        "",
        "yes",
        "REQ-DA-203",
    ],
]

HOP2_FIELDMAP = [
    # v_dim_customer_src -> dim_customer (spec rows 3-5,7-15,17 + row 18 join)
    _fm(
        "dim_customer_migration",
        "cunbr",
        "involved_party_key",
        "string",
        "no",
        "60",
        "",
        "concat('CBS', 'CUST', cubk, cunbr)",
        desc="party key = 'CBS' + 'CUST' + bank number + customer number",
        req="REQ-DA-205",
    ),
    _fm("dim_customer_migration", "load_date", "load_date", "date", "no", desc="direct copy"),
    _fm(
        "dim_customer_migration",
        "cupers",
        "customer_type_cd",
        "string",
        "no",
        "1",
        "",
        "casewhen(trim(cupers), 'N', 'C', trim(cupers))",
        desc="recode customer type 'N' to 'C'; otherwise keep the trimmed code",
        req="REQ-DA-206",
    ),
    _fm(
        "dim_customer_migration",
        "culegenty",
        "legal_entity_cd",
        "string",
        "yes",
        "50",
        desc="direct copy",
    ),
    _fm(
        "dim_customer_migration",
        "cubk",
        "bank_num",
        "string",
        "no",
        "3",
        "",
        "lpad(trim(cubk), 3, '0')",
        desc="bank number, zero-padded to 3 digits",
        req="REQ-DA-207",
    ),
    _fm(
        "dim_customer_migration",
        "cunbr",
        "appl_key",
        "string",
        "no",
        "60",
        "",
        "literal('CBS')",
        desc="constant 'CBS' — every row carries the CBS application key",
        req="REQ-DA-208",
    ),
    _fm(
        "dim_customer_migration",
        "cusic",
        "naics_cd",
        "string",
        "yes",
        "6",
        "",
        "lpad(cusic, 6, '0')",
        desc="NAICS industry code, zero-padded to 6 digits",
        req="REQ-DA-209",
    ),
    _fm(
        "dim_customer_migration",
        "cuna1",
        "full_nm",
        "string",
        "no",
        "256",
        "",
        "trim(cuna1)",
        "yes",
        "customer name, whitespace trimmed",
        req="REQ-DA-210",
    ),
    _fm(
        "dim_customer_migration",
        "cucfnm",
        "first_nm",
        "string",
        "yes",
        "256",
        "",
        "casewhen(cupers, 'P', cucfnm, None)",
        "yes",
        "first name kept for persons ('P') only; NULL otherwise",
        req="REQ-DA-211",
    ),
    _fm(
        "dim_customer_migration",
        "cussnr",
        "tax_id_number",
        "string",
        "yes",
        "9",
        "",
        "lpad(cussnr, 9, '0')",
        "yes",
        "tax id, zero-padded to 9 digits",
        req="REQ-DA-212",
    ),
    _fm(
        "dim_customer_migration",
        "cussty",
        "tax_id_number_type_cd",
        "string",
        "no",
        "1",
        "",
        "casewhen(ifnull(cussty, ''), '', 'S', cussty)",
        desc="blank tax-id type defaults to 'S' (SSN)",
        req="REQ-DA-213",
    ),
    _fm(
        "dim_customer_migration",
        "cujbdt",
        "birth_dt",
        "date",
        "yes",
        "",
        "",
        "date_fmt(cujbdt, '%Y-%m-%d', '%Y%m%d')",
        "yes",
        "birth date reformatted from YYYYMMDD to YYYY-MM-DD",
        req="REQ-DA-214",
    ),
    _fm(
        "dim_customer_migration",
        "cuclnm",
        "last_nm",
        "string",
        "yes",
        "256",
        "",
        "casewhen(cupers, 'P', cuclnm, None)",
        "yes",
        "last name kept for persons ('P') only; NULL otherwise",
        req="REQ-DA-211",
    ),
    _fm(
        "dim_customer_migration",
        "attribute_value_desc",
        "tax_id_number_type_desc",
        "string",
        "yes",
        "256",
        desc="tax-id-type description carried from the join view (DIM_REFERENCE decode)",
        req="REQ-DA-215",
    ),
    # cup00301 -> dim_customer_address (spec row 6 + carried address columns)
    _fm(
        "dim_customer_address_migration",
        "cunbr",
        "involved_party_key",
        "string",
        "no",
        "60",
        "",
        "concat('CBS', 'CUST', cubk, cunbr)",
        desc="party key = 'CBS' + 'CUST' + bank number + customer number",
        req="REQ-DA-205",
    ),
    _fm(
        "dim_customer_address_migration",
        "cucncd",
        "address_country_cd",
        "string",
        "no",
        "5",
        "",
        "lpad(cucncd, 5, '0')",
        desc="country code, zero-padded to 5 digits",
        req="REQ-DA-216",
    ),
    _fm(
        "dim_customer_address_migration",
        "cuadl1",
        "address_line_1",
        "string",
        "yes",
        "80",
        "",
        "trim(cuadl1)",
        "yes",
        "address line 1, whitespace trimmed",
        req="REQ-DA-217",
    ),
    _fm(
        "dim_customer_address_migration",
        "cuadl2",
        "address_line_2",
        "string",
        "yes",
        "80",
        "",
        "trim(cuadl2)",
        "yes",
        "address line 2, whitespace trimmed",
        req="REQ-DA-218",
    ),
    _fm(
        "dim_customer_address_migration",
        "cuadct",
        "city",
        "string",
        "yes",
        "60",
        "",
        "",
        "yes",
        "direct copy",
    ),
    _fm(
        "dim_customer_address_migration",
        "cuadst",
        "state_cd",
        "string",
        "yes",
        "30",
        desc="direct copy",
    ),
    _fm(
        "dim_customer_address_migration",
        "cuadpc",
        "postal_cd",
        "string",
        "yes",
        "15",
        desc="direct copy",
    ),
    _fm(
        "dim_customer_address_migration",
        "load_date",
        "load_date",
        "date",
        "no",
        desc="direct copy",
    ),
    # dim_reference -> dim_reference (reference data lands in IL unchanged)
    _fm(
        "dim_reference_migration",
        "attribute_cd",
        "attribute_cd",
        "string",
        "no",
        "50",
        desc="direct copy",
    ),
    _fm(
        "dim_reference_migration",
        "attribute_value_desc",
        "attribute_value_desc",
        "string",
        "no",
        "256",
        desc="direct copy",
    ),
    _fm("dim_reference_migration", "load_date", "load_date", "date", "no", desc="direct copy"),
    _fm(
        "dim_reference_migration",
        "attribute_map_cd",
        "attribute_map_cd",
        "string",
        "yes",
        "50",
        desc="direct copy",
    ),
]

# Post-transform format check on the integration layer — cujbdt lands as
# YYYYMMDD on hop-1 (covered by BR-JOBDATE-FORMAT); the hop-2 transform
# date_fmt(cujbdt,'%Y-%m-%d','%Y%m%d') must produce ISO YYYY-MM-DD birth_dt.
HOP2_RULES = [
    [
        "BR-BIRTHDT-ISO",
        "DATE_FORMAT",
        "",
        "",
        "dim_customer",
        "birth_dt",
        "",
        "%Y-%m-%d",
        "MEDIUM",
        "",
        "involved_party_key",
        "yes",
        "birth date must be ISO YYYY-MM-DD after the hop-2 reformat",
        "REQ-DI-202",
    ],
]

HOP2_CONTROLS = [
    [
        "dim_customer_migration",
        "dim_customer_rows",
        "count",
        "",
        "",
        "dim_customer_rows",
        "yes",
        "REQ-DC-201",
    ],
    [
        "dim_customer_migration",
        "dim_customer_by_type",
        "count",
        "",
        "customer_type_cd",
        "dim_customer_by_type",
        "yes",
        "REQ-DA-204",
    ],
    [
        "dim_customer_address_migration",
        "dim_address_rows",
        "count",
        "",
        "",
        "dim_address_rows",
        "yes",
        "REQ-DC-202",
    ],
    [
        "dim_reference_migration",
        "dim_ref_rows",
        "count",
        "",
        "",
        "dim_ref_rows",
        "yes",
        "REQ-DC-203",
    ],
]

HOP2_RELATIONS = [
    [
        "dim_customer_migration",
        "involved_party_key",
        "dim_customer_address_migration",
        "involved_party_key",
        "yes",
        "REQ-DI-201",
    ],
]

HOP2_README = [
    ["Phase-1 hop 2: signature -> wdr_il_core (integration layer)"],
    [""],
    ["dim_customer source is SIGNATURE.V_DIM_CUSTOMER_SRC — a read-only"],
    ["view joining CUP00301 + CUP01201 (legal entity) + DIM_REFERENCE"],
    ["(tax-id-type description). Lookup semantics per spec row 18."],
    [""],
    ["Spec resolutions:"],
    ["  - APPL_KEY: hardcoded 'CBS' (row 9 authoritative); CUSIC row"],
    ["    contributes NAICS_CD only."],
    ["  - CASE arms without ELSE (CUPERS first/last name): ELSE NULL"],
    ["    per spec — nullable target columns carry it."],
    ["  - DIM_REFERENCE also lands ATTRIBUTE_MAP_CD — the NAICS lookup key"],
    ["    used by hop-3 CUSTOMER_CATEGORY."],
]


# ---------------------------------------------------------------- hop 3
HOP3_TABLES = [
    [
        "cl_party_snpsht_migration",
        "wdr_il_core",
        "v_party_snpsht_src",
        "wdr_cl_fds_core",
        "cl_party_snpsht",
        "party_identifier",
        "",
        "full",
        "",
        "yes",
        "REQ-DA-301",
    ],
]

HOP3_FIELDMAP = [
    _fm(
        "cl_party_snpsht_migration",
        "involved_party_key",
        "party_identifier",
        "string",
        "no",
        "60",
        desc="direct copy",
    ),
    _fm("cl_party_snpsht_migration", "load_date", "as_of_date", "date", "no", desc="direct copy"),
    _fm(
        "cl_party_snpsht_migration",
        "customer_type_cd",
        "business_domain",
        "string",
        "no",
        "100",
        "",
        "casewhen(customer_type_cd, 'P', 'Personal', 'N', 'Non-Personal', "
        "'C', 'Commercial', 'B', 'Business', 'UNSPECIFIED')",
        desc="decode customer type: 'P'->Personal, 'N'->Non-Personal, 'C'->Commercial, "
        "'B'->Business, otherwise UNSPECIFIED",
        req="REQ-DA-303",
    ),
    _fm(
        "cl_party_snpsht_migration",
        "address_country_cd",
        "domicile_country_code",
        "string",
        "no",
        "20",
        "",
        "casewhen(regexp_match(address_country_cd, '^[A-Za-z]{2}$'), 'True', "
        "address_country_cd, 'US')",
        desc="keep the country code only if it is two letters, otherwise 'US' — padded "
        "numeric codes never match, so 'US' always wins (per spec)",
        req="REQ-DA-304",
    ),
    _fm(
        "cl_party_snpsht_migration",
        "legal_entity_cd",
        "organization_constituent_type_code",
        "string",
        "yes",
        "20",
        desc="direct copy",
    ),
    _fm(
        "cl_party_snpsht_migration",
        "bank_num",
        "bank_num",
        "string",
        "no",
        "50",
        desc="direct copy",
    ),
    _fm(
        "cl_party_snpsht_migration",
        "appl_key",
        "data_source_code",
        "string",
        "no",
        "20",
        desc="direct copy",
    ),
    _fm(
        "cl_party_snpsht_migration",
        "appl_key",
        "sor_party_type",
        "string",
        "yes",
        "50",
        "",
        "casewhen(appl_key, 'PAR', 'BROKER-DEALER', naics_cd)",
        desc="'PAR' maps to 'BROKER-DEALER'; otherwise the NAICS code",
        req="REQ-DA-305",
    ),
    # SPEC GAP (row 11): "Complex Logic - details shared in next tab" — the
    # tab does not exist. Sentinel keeps the column verifiably present and
    # makes the gap visible in data + evidence instead of silently dropping it.
    _fm(
        "cl_party_snpsht_migration",
        "involved_party_key",
        "party_type",
        "string",
        "no",
        "60",
        "",
        "literal('PENDING_SPEC')",
        desc="spec gap sentinel — party type rule pending client spec",
        req="REQ-DA-306",
    ),
    _fm(
        "cl_party_snpsht_migration",
        "full_nm",
        "party_full_name",
        "string",
        "no",
        "255",
        "",
        "",
        "yes",
        "direct copy",
    ),
    _fm(
        "cl_party_snpsht_migration",
        "first_nm",
        "party_name",
        "string",
        "yes",
        "255",
        "",
        "",
        "yes",
        "direct copy",
    ),
    _fm(
        "cl_party_snpsht_migration",
        "last_nm",
        "party_last_name",
        "string",
        "yes",
        "255",
        desc="direct copy",
    ),
    _fm(
        "cl_party_snpsht_migration",
        "tax_id_number",
        "party_tax_identifier",
        "string",
        "yes",
        "60",
        "",
        "",
        "yes",
        "direct copy",
    ),
    _fm(
        "cl_party_snpsht_migration",
        "tax_id_number_type_cd",
        "party_tax_identifier_format_indicator",
        "string",
        "no",
        "1",
        "",
        "substr(tax_id_number_type_cd, 0, 1)",
        desc="first character of the tax-id type code",
        req="REQ-DA-307",
    ),
    _fm(
        "cl_party_snpsht_migration",
        "birth_dt",
        "date_of_birth",
        "date",
        "yes",
        "",
        "",
        "",
        "yes",
        "direct copy",
    ),
    _fm(
        "cl_party_snpsht_migration",
        "tax_id_number_type_desc",
        "wbs_party_tax_identifier_description",
        "string",
        "yes",
        "255",
        desc="direct copy",
    ),
    # CUSTOMER_CATEGORY (client CASE): DIM_REFERENCE lookup on naics_cd wins;
    # else searched CASE on customer_type_cd + full_nm markers. NULL for
    # unmatched 'C'/'B' names — faithful to the spec snippet's missing ELSE.
    # Last FieldMap row so contract column order matches the appended column
    # on pre-existing targets.
    _fm(
        "cl_party_snpsht_migration",
        "attribute_cd",
        "customer_category",
        "string",
        "yes",
        "50",
        "",
        CUSTOMER_CATEGORY_TRANSFORM,
        desc="NAICS-match in DIM_REFERENCE wins; else 'P'->INDIVIDUAL, 'C'/'B' + LLC/INC "
        "in name->CORP_SMALL, +BANK->BANK, else NULL",
        req="REQ-DA-308",
    ),
]

HOP3_CONTROLS = [
    [
        "cl_party_snpsht_migration",
        "cl_party_rows",
        "count",
        "",
        "",
        "cl_party_rows",
        "yes",
        "REQ-DC-301",
    ],
    [
        "cl_party_snpsht_migration",
        "cl_party_by_domain",
        "count",
        "",
        "business_domain",
        "cl_party_by_domain",
        "yes",
        "REQ-DA-302",
    ],
]

HOP3_README = [
    ["Phase-1 hop 3: wdr_il_core -> wdr_cl_fds_core (consumption layer)"],
    [""],
    ["Source is WDR_IL_CORE.V_PARTY_SNPSHT_SRC — read-only view joining"],
    ["DIM_CUSTOMER LEFT JOIN DIM_CUSTOMER_ADDRESS on involved_party_key."],
    [""],
    ["Spec resolutions:"],
    ["  - BUSINESS_DOMAIN CASE has no ELSE in spec -> ELSE 'UNSPECIFIED'."],
    ["  - DOMICILE_COUNTRY_CODE: REGEXP_LIKE '^[A-Za-z]{2}$' on the padded"],
    ["    numeric ADDRESS_COUNTRY_CD always fails -> 'US' (faithful to spec)."],
    ["  - PARTY_TYPE: 'Complex Logic - details shared in next tab' does not"],
    ["    exist -> literal('PENDING_SPEC') sentinel until client supplies it."],
    ["  - PARTY_TAX_IDENTIFIER_FORMAT_INDICATOR: LEFT(x,1) -> substr(x,0,1)."],
    ["  - CUSTOMER_CATEGORY: client CASE — DIM_REFERENCE match on"],
    ["    ATTRIBUTE_MAP_CD = NAICS_CD returns ATTRIBUTE_CD; else 'P' ->"],
    ["    INDIVIDUAL, 'C'/'B' + name LIKE %LLC%/%INC% -> CORP_SMALL,"],
    ["    'C'/'B' + name LIKE %BANK% -> BANK, otherwise NULL (spec has"],
    ["    no ELSE — declared ELSE NULL). V_PARTY_SNPSHT_SRC LEFT JOINs"],
    ["    DIM_REFERENCE to expose attribute_cd."],
]


def _sheet(wb, title, headers, rows, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    if first:
        ws.title = title
    if headers:
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    return ws


def _write(name, tables, fieldmap, controls, relations, readme, rules=None):
    wb = Workbook()
    _sheet(wb, "README", None, readme, first=True)
    wb["README"].column_dimensions["A"].width = 90
    _sheet(wb, "Tables", TABLES_HEADERS, tables)
    _sheet(wb, "FieldMap", FIELDMAP_HEADERS, fieldmap)
    _sheet(wb, "Controls", CONTROLS_HEADERS, controls)
    _sheet(wb, "Relations", RELATIONS_HEADERS, relations)
    _sheet(wb, "Mapping", MAPPING_HEADERS, [])
    if rules:
        rs = _sheet(wb, "Rules", RULES_HEADERS, rules)
        rs.column_dimensions["B"].width = 24
        rs.column_dimensions["G"].width = 70
        rs.column_dimensions["M"].width = 60
    out = OUT_DIR / name
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}")


def main():
    _write(
        "phase1_hop1_csv_to_stage.xlsx",
        HOP1_TABLES,
        HOP1_FIELDMAP,
        HOP1_CONTROLS,
        HOP1_RELATIONS,
        HOP1_README,
        HOP1_RULES,
    )
    _write(
        "phase1_hop2_stage_to_il.xlsx",
        HOP2_TABLES,
        HOP2_FIELDMAP,
        HOP2_CONTROLS,
        HOP2_RELATIONS,
        HOP2_README,
        HOP2_RULES,
    )
    _write(
        "phase1_hop3_il_to_consumption.xlsx",
        HOP3_TABLES,
        HOP3_FIELDMAP,
        HOP3_CONTROLS,
        [],
        HOP3_README,
    )


if __name__ == "__main__":
    main()
