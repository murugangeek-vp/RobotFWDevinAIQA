"""Build data/mapping/bank1_to_mss.xlsx — the Bank-1 (ent_* schema) -> MSS
Banking (mss_* schema) migration workbook.

Same sheet contract as bank_mapping.xlsx: Tables / FieldMap / Mapping /
Relations / Controls. Demonstrates the realistic cases a real bank
migration hits: renames, code translation, numeric banding (risk score ->
classification), boolean->enum, a WHERE scope filter, PII flags, cross-
mapping FK relations, certified trailer controls — and honest gaps (target
columns with no source stay unmapped; unmapped source tables show as
insights unless run.fail_on_unmapped is set).

Run: .venv/Scripts/python.exe scripts/generate_bank1_workbook.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

REPO = Path(__file__).resolve().parent.parent
OUT = REPO / "data" / "mapping" / "bank1_to_mss.xlsx"

TABLES_HEADERS = [
    "Mapping ID", "Source Schema", "Source Table",
    "Target Schema", "Target Table", "Key Columns",
    "Where Clause", "Recon Mode", "Batch Size", "Enabled",
]

# transaction_journal, ent_countries/states/cities have no MSS target —
# they stay unmapped and surface as evidence insights (fail_on_unmapped
# in the env turns that into a hard FAIL if the program requires it).
TABLES = [
    ["customer_migration", "bank1", "customer_master", "PUBLIC",
     "mss_customer_master", "customer_urn", "", "full", "", "yes"],
    ["product_migration", "bank1", "bank_products", "PUBLIC",
     "mss_product_master", "product_code", "", "full", "", "yes"],
    ["branch_migration", "bank1", "ent_branches", "PUBLIC",
     "mss_branches", "branch_id", "is_active = TRUE", "full", "", "yes"],
    ["account_migration", "bank1", "account_master", "PUBLIC",
     "mss_accounts", "account_number", "", "full", "", "yes"],
    ["savings_ledger_migration", "bank1", "ledger_savings", "PUBLIC",
     "mss_ledger_savings", "account_number", "", "full", "", "yes"],
    ["fd_ledger_migration", "bank1", "ledger_fixed_deposits", "PUBLIC",
     "mss_ledger_fixed_deposits", "account_number", "", "full", "", "yes"],
    ["loan_ledger_migration", "bank1", "ledger_loans", "PUBLIC",
     "mss_ledger_loans", "account_number", "", "full", "", "yes"],
    ["role_migration", "bank1", "app_roles", "PUBLIC",
     "mss_staff_roles", "role_id", "", "full", "", "yes"],
    ["staff_migration", "bank1", "app_users", "PUBLIC",
     "mss_staff_users", "staff_emp_id", "", "full", "", "yes"],
]

FIELDMAP_HEADERS = [
    "Mapping ID", "Source Column", "Target Column", "Target Type",
    "Nullable", "Max Length", "Scale", "Transformation", "PII", "Enabled",
]

FIELDMAP = [
    # customer_master -> mss_customer_master
    ["customer_migration", "customer_urn", "customer_urn", "string", "no", "30", "", "", "no", "yes"],
    ["customer_migration", "legal_name", "full_legal_name", "string", "no", "250", "", "strip({legal_name})", "yes", "yes"],
    ["customer_migration", "customer_type", "customer_segment", "string", "no", "15", "", "map(customer_type, segment_codes)", "no", "yes"],
    ["customer_migration", "risk_rating_score", "risk_classification", "string", "no", "10", "", "band(risk_rating_score, 40, 70, 'LOW', 'MEDIUM', 'HIGH')", "no", "yes"],
    ["customer_migration", "aml_kyc_status", "kyc_status", "string", "no", "15", "", "map(aml_kyc_status, kyc_codes)", "no", "yes"],
    # entity_id / date_of_boarding: no source column — documented gap.
    # kyc_expiry_date: semantically different from date_of_boarding — dropped.
    # bank_products -> mss_product_master
    ["product_migration", "product_code", "product_code", "string", "no", "30", "", "", "no", "yes"],
    ["product_migration", "product_category", "category_id", "string", "no", "15", "", "map(product_category, cat_codes)", "no", "yes"],
    ["product_migration", "product_display_name", "product_name", "string", "no", "100", "", "trim({product_display_name})", "no", "yes"],
    # base_interest_rate moves to mss_product_interest_rules (not mapped here);
    # entity_id / currency / is_available have no source — gaps.
    # ent_branches -> mss_branches
    ["branch_migration", "branch_routing_id", "branch_id", "string", "no", "20", "", "", "no", "yes"],
    ["branch_migration", "branch_routing_id", "clearing_code_value", "string", "no", "20", "", "", "no", "yes"],
    ["branch_migration", "branch_name", "branch_name", "string", "no", "150", "", "trim({branch_name})", "no", "yes"],
    ["branch_migration", "branch_tier", "branch_tier", "string", "no", "10", "", "", "no", "yes"],
    # entity_id / region_id / clearing_code_type: no source — gaps.
    # account_master -> mss_accounts
    ["account_migration", "account_number", "account_number", "string", "no", "34", "", "", "no", "yes"],
    ["account_migration", "customer_urn", "customer_urn", "string", "no", "30", "", "", "no", "yes"],
    ["account_migration", "product_code", "product_code", "string", "no", "30", "", "", "no", "yes"],
    ["account_migration", "hosting_branch_id", "branch_id", "string", "no", "20", "", "", "no", "yes"],
    ["account_migration", "account_status", "account_status", "string", "no", "15", "", "", "no", "yes"],
    ["account_migration", "is_insider_account", "is_insider_account", "boolean", "no", "", "", "", "no", "yes"],
    ["account_migration", "opened_timestamp", "opened_date", "timestamp", "no", "", "", "", "no", "yes"],
    # currency_iso: no target column on mss_accounts (currency lives on the
    # product) — documented as dropped.
    # ledger_savings -> mss_ledger_savings
    ["savings_ledger_migration", "account_number", "account_number", "string", "no", "34", "", "", "no", "yes"],
    ["savings_ledger_migration", "available_balance", "available_balance", "decimal", "no", "", "4", "round({available_balance}, 4)", "no", "yes"],
    ["savings_ledger_migration", "book_balance", "ledger_balance", "decimal", "no", "", "4", "round({book_balance}, 4)", "no", "yes"],
    # overdraft_limit: no target column — dropped.
    # ledger_fixed_deposits -> mss_ledger_fixed_deposits
    ["fd_ledger_migration", "account_number", "account_number", "string", "no", "34", "", "", "no", "yes"],
    ["fd_ledger_migration", "principal_deposit", "principal_amount", "decimal", "no", "", "4", "round({principal_deposit}, 4)", "no", "yes"],
    ["fd_ledger_migration", "maturity_amount", "maturity_amount", "decimal", "no", "", "4", "round({maturity_amount}, 4)", "no", "yes"],
    ["fd_ledger_migration", "maturity_date", "maturity_date", "date", "no", "", "", "", "no", "yes"],
    ["fd_ledger_migration", "auto_renewal_flag", "auto_renewal_option", "string", "no", "30", "", "casewhen(auto_renewal_flag, 'True', 'PRINCIPAL_AND_INTEREST', 'False', 'NONE')", "no", "yes"],
    # tenure_days / interest_payout_account: no source — gaps.
    # ledger_loans -> mss_ledger_loans
    ["loan_ledger_migration", "account_number", "account_number", "string", "no", "34", "", "", "no", "yes"],
    ["loan_ledger_migration", "sanctioned_amount", "sanctioned_limit", "decimal", "no", "", "4", "round({sanctioned_amount}, 4)", "no", "yes"],
    ["loan_ledger_migration", "principal_outstanding", "principal_outstanding", "decimal", "no", "", "4", "round({principal_outstanding}, 4)", "no", "yes"],
    ["loan_ledger_migration", "emi_amount", "emi_amount", "decimal", "no", "", "4", "round({emi_amount}, 4)", "no", "yes"],
    ["loan_ledger_migration", "tenure_months", "tenure_months", "integer", "no", "", "", "", "no", "yes"],
    ["loan_ledger_migration", "next_due_date", "next_emi_date", "date", "no", "", "", "", "no", "yes"],
    # interest_outstanding / asset_classification: no source — gaps.
    # app_roles -> mss_staff_roles
    ["role_migration", "role_id", "role_id", "string", "no", "30", "", "map(role_id, role_codes)", "no", "yes"],
    # role_clearance_level: no source — gap.
    # app_users -> mss_staff_users
    ["staff_migration", "user_emp_id", "staff_emp_id", "string", "no", "20", "", "", "no", "yes"],
    ["staff_migration", "username", "username", "string", "no", "50", "", "", "yes", "yes"],
    ["staff_migration", "primary_role_id", "role_id", "string", "no", "30", "", "map(primary_role_id, role_codes)", "no", "yes"],
    ["staff_migration", "home_branch_routing_id", "assigned_branch_id", "string", "no", "20", "", "", "no", "yes"],
    # security_clearance_band / ip_whitelist_range / last_login_timestamp:
    # security attributes with no MSS columns — dropped; is_active unmapped.
]

MAPPING_HEADERS = ["Mapping Name", "Code", "Value"]
MAPPING = [
    ["segment_codes", "RETAIL", "RETAIL"],
    ["segment_codes", "CORPORATE", "CORPORATE"],
    ["segment_codes", "SME", "SME"],
    ["segment_codes", "HNI", "HNI"],
    ["kyc_codes", "VERIFIED", "VERIFIED"],
    ["kyc_codes", "PENDING", "PENDING"],
    ["kyc_codes", "SUSPENDED", "REJECTED"],
    ["cat_codes", "SAVINGS", "SAVINGS"],
    ["cat_codes", "LOAN", "LOAN"],
    ["cat_codes", "FD", "FIXED_DEPOSIT"],
    ["role_codes", "CORE_TELLER", "TELLER"],
    ["role_codes", "CREDIT_ANALYST", "CREDIT_OFFICER"],
    ["role_codes", "CROSS_BORDER_AUDITOR", "AUDITOR"],
    ["role_codes", "BRANCH_MANAGER", "BRANCH_MANAGER"],
]

RELATIONS_HEADERS = [
    "Parent Mapping", "Parent Column", "Child Mapping", "Child Column", "Enabled",
]
RELATIONS = [
    ["customer_migration", "customer_urn", "account_migration", "customer_urn", "yes"],
    ["product_migration", "product_code", "account_migration", "product_code", "yes"],
    ["branch_migration", "branch_id", "account_migration", "branch_id", "yes"],
    ["account_migration", "account_number", "savings_ledger_migration", "account_number", "yes"],
    ["account_migration", "account_number", "fd_ledger_migration", "account_number", "yes"],
    ["account_migration", "account_number", "loan_ledger_migration", "account_number", "yes"],
    ["role_migration", "role_id", "staff_migration", "role_id", "yes"],
]

CONTROLS_HEADERS = [
    "Mapping ID", "Control ID", "Type", "Column", "Group By", "File Name", "Enabled",
]
CONTROLS = [
    ["customer_migration", "customer_count", "count", "", "", "customers", "yes"],
    ["account_migration", "account_rows", "count", "", "", "accounts", "yes"],
    ["account_migration", "accounts_by_status", "count", "", "account_status", "account_status", "yes"],
    ["savings_ledger_migration", "savings_book_total", "sum", "ledger_balance", "", "savings_book", "yes"],
    ["savings_ledger_migration", "savings_rows", "count", "", "", "savings_rows", "yes"],
    ["loan_ledger_migration", "loans_principal_total", "sum", "principal_outstanding", "", "loans_principal", "yes"],
    ["fd_ledger_migration", "fd_principal_total", "sum", "principal_amount", "", "fd_principal", "yes"],
]

README = [
    ["Bank-1 -> MSS Banking migration mapping workbook"],
    [""],
    ["Demonstrates a real bank migration: enterprise ent_* source schema"],
    ["mapped onto the MSS mss_* core-banking model."],
    [""],
    ["Known gaps (documented, not hidden):"],
    ["  - transaction_journal + ent_countries/states/cities have no MSS"],
    ["    target tables — they surface as unmapped-source insights."],
    ["  - mss_* target columns with no source stay unmapped (entity_id,"],
    ["    currency, kyc boarding date, clearance levels, etc.)."],
    ["  - account_master.currency_iso is dropped — currency is on the"],
    ["    product in the MSS model."],
    [""],
    ["Connection details are NOT in this workbook — they live in"],
    ["config/environments/bank1_to_mss.yaml; credentials are env vars only."],
]


def _sheet(wb, title, headers, rows, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    if first:
        ws.title = title
    if headers:
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    return ws


def main():
    wb = Workbook()
    _sheet(wb, "README", None, README, first=True)
    wb["README"].column_dimensions["A"].width = 90
    _sheet(wb, "Tables", TABLES_HEADERS, TABLES)
    _sheet(wb, "FieldMap", FIELDMAP_HEADERS, FIELDMAP)
    _sheet(wb, "Controls", CONTROLS_HEADERS, CONTROLS)
    _sheet(wb, "Relations", RELATIONS_HEADERS, RELATIONS)
    _sheet(wb, "Mapping", MAPPING_HEADERS, MAPPING)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
