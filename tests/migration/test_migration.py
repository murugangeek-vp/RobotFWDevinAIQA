import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
import pandas as pd
from test_streaming import contract, frames

from libs.adapters.postgres_source import PostgresSource
from libs.adapters.snowflake_source import SnowflakeSource
from libs.adapters.snowflake_target import SnowflakeTarget
from libs.adapters.staging import ordered_chunks
from libs.engine import controls as engine_controls
from libs.engine import relations as engine_relations
from libs.engine.mapping_transforms import evaluate, mapped_rows, parse_expression
from libs.engine.reconcile import compare_chunked, normalize
from libs.mapping import builder, parser
from libs.reporting.report import write_evidence
from libs.robot.MigrationLibrary import MigrationLibrary
from libs.robot.ReconciliationLibrary import ReconciliationLibrary

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "data/mapping/phase1_hop2_stage_to_int.xlsx"


class MappingTests(unittest.TestCase):
    def setUp(self):
        self.tables, self.codes = parser.load_workbook(str(WORKBOOK))

    def compile(self, table=None):
        return builder.build_contract(
            table or self.tables[0],
            self.codes,
            {"type": "postgres", "host": "localhost", "database": "synthetic"},
            {"type": "snowflake"},
        )

    def test_all_enabled_mappings_compile(self):
        enabled = [t for t in self.tables if t.enabled]
        self.assertEqual(len(enabled), 3)
        for t in enabled:
            self.assertEqual(self.compile(t).name, t.mapping_id)

    def test_composite_key_does_not_require_individually_unique_columns(self):
        table = copy.deepcopy(self.tables[0])
        table.keys = ["involved_party_key", "load_date"]
        self.assertFalse(self.compile(table).column("load_date")["unique"])

    def test_band_transform_buckets_scores(self):
        from libs.engine.reconcile import apply_transform

        row = pd.Series({"risk_rating_score": 55})
        expr = "band(risk_rating_score, 40, 70, 'LOW', 'MEDIUM', 'HIGH')"
        self.assertEqual(apply_transform(expr, row), "MEDIUM")
        for score, want in (
            (1, "LOW"),
            (40, "LOW"),
            (41, "MEDIUM"),
            (70, "MEDIUM"),
            (71, "HIGH"),
            (100, "HIGH"),
        ):
            self.assertEqual(apply_transform(expr, pd.Series({"risk_rating_score": score})), want)

    def test_unknown_transform_rejected(self):
        for expression in (
            "evil(email)",
            "upper(email).__class__",
            "__import__('os')",
            "upper(email, **x)",
        ):
            with self.subTest(expression=expression), self.assertRaises((ValueError, SyntaxError)):
                parse_expression(expression)

    def test_transform_preserves_commas(self):
        self.assertEqual(
            evaluate(parse_expression("upper({name})"), {"name": "Smith, Jane"}, {}), "SMITH, JANE"
        )

    def test_null_stays_null(self):
        self.assertIsNone(evaluate(parse_expression("upper({name})"), {"name": None}, {}))

    def test_unmapped_code_fails(self):
        with self.assertRaisesRegex(ValueError, "Unmapped"):
            evaluate(
                parse_expression("map(code, status)"), {"code": "bad"}, {"status": {"A": "active"}}
            )

    def test_explicit_date_format(self):
        self.assertEqual(
            evaluate(
                parse_expression("date_fmt(dob, '%Y-%m-%d', '%m/%d/%Y')"), {"dob": "12/31/1990"}, {}
            ),
            "1990-12-31",
        )

    def test_strict_string_comparison(self):
        spec = {"type": "string", "exact": True}
        self.assertNotEqual(normalize(" a ", spec), normalize("a", spec))
        self.assertNotEqual(normalize("", spec), normalize(None, spec))

    def test_empty_transform_frame(self):
        c = self.compile()
        self.assertEqual(
            list(mapped_rows(pd.DataFrame(columns=c.source_columns), c).columns),
            [col["name"] for col in c.columns],
        )

    def test_invalid_workbook_cells(self):
        cases = [
            ("Tables", "J2", "maybe"),
            ("Tables", "I2", 12.5),
            ("FieldMap", "D2", "unknown"),
            ("FieldMap", "H2", "=NOW()"),
            ("FieldMap", "E2", None),
            ("Tables", "A1", "broken"),
        ]
        for sheet, cell, value in cases:
            with self.subTest(sheet=sheet, cell=cell), tempfile.TemporaryDirectory() as tmp:
                wb = openpyxl.load_workbook(WORKBOOK)
                wb[sheet][cell] = value
                path = Path(tmp) / "bad.xlsx"
                wb.save(path)
                wb.close()
                with self.assertRaises(ValueError):
                    tables, codes = parser.load_workbook(str(path))
                    builder.build_contract(
                        tables[0], codes, {"type": "postgres"}, {"type": "snowflake"}
                    )

    def test_disabled_mapping_cannot_be_selected(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.load_workbook(WORKBOOK)
            wb["Tables"]["J2"] = "no"  # Enabled column
            path = Path(tmp) / "disabled.xlsx"
            wb.save(path)
            wb.close()
            tables, _ = parser.load_workbook(str(path))
            disabled = next(t for t in tables if not t.enabled)
            self.assertEqual(disabled.mapping_id, tables[0].mapping_id)


class StagingTests(unittest.TestCase):
    def test_staging_cleans_up_after_failure(self):
        c = contract()
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "null"):
                list(ordered_chunks(frames([(None, "a", "b")]), c, 2, tmp))
            self.assertEqual(list(Path(tmp).iterdir()), [])

    def test_external_sort(self):
        from contextlib import closing

        c = contract()
        with tempfile.TemporaryDirectory() as tmp:
            with (
                closing(
                    ordered_chunks(
                        frames([(20, "a", "b"), (1, "a", "b"), (2, "a", "b")]), c, 2, tmp
                    )
                ) as source,
                closing(
                    ordered_chunks(
                        frames([(2, "a", "b"), (20, "a", "b"), (1, "a", "b")]), c, 1, tmp
                    )
                ) as target,
            ):
                self.assertEqual(compare_chunked(source, target, c).mismatch_count, 0)
            self.assertEqual(list(Path(tmp).iterdir()), [])


class SafetyAndEvidenceTests(unittest.TestCase):
    def test_query_rejects_mutations_before_execution(self):
        adapter = SnowflakeTarget.__new__(SnowflakeTarget)
        adapter.conn = MagicMock()
        for statement in (
            "DELETE FROM test",
            "SELECT 1; DROP TABLE test",
            "SELECT * INTO x FROM test",
        ):
            with self.subTest(statement=statement), self.assertRaises(ValueError):
                adapter.query(statement)
        adapter.conn.cursor.assert_not_called()

    def test_earlier_failures_remain_failures(self):
        with tempfile.TemporaryDirectory() as tmp:
            results = [
                {"mapping": "a", "stage": "recon", "status": "FAIL", "mismatches": 1},
                {"mapping": "b", "stage": "recon", "status": "PASS"},
            ]
            write_evidence(results, {"selected": ["a", "b"]}, tmp)
            data = json.loads((Path(tmp) / "run_summary.json").read_text())
            self.assertEqual(data["verdict"], "FAIL")
            self.assertEqual(data["coverage"], "PARTIAL")
            self.assertFalse(data["production_signoff"])

    def test_empty_run_cannot_pass(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(write_evidence([], {}, tmp), "NOT_RUN")

    def test_excel_formula_injection_disabled(self):
        with tempfile.TemporaryDirectory() as tmp:
            write_evidence([{"mapping": "=1+1", "stage": "recon", "status": "PASS"}], {}, tmp)
            wb = openpyxl.load_workbook(Path(tmp) / "migration_summary.xlsx")
            self.assertEqual(wb["Stage Results"]["A2"].data_type, "s")
            wb.close()


def _library(env):
    lib = MigrationLibrary.__new__(MigrationLibrary)
    lib.env = copy.deepcopy(env)
    lib.meta = {}
    lib._snapshot_holder = None
    return lib


class SnapshotTests(unittest.TestCase):
    ENV = {
        "source": {
            "type": "postgres",
            "host": "h",
            "port": 5432,
            "database": "d",
            "user": "u",
            "password": "p",
        },
        "target": {"type": "snowflake", "fake": True},
        "run": {"snapshot_id": "cut-2026-01-01"},
    }

    def test_malformed_exported_snapshot_rejected_before_connect(self):
        for bad in ("0001; DROP TABLE t", "snap 1", "'x'", ""):
            with (
                self.subTest(snapshot=bad),
                self.assertRaises(ValueError),
                patch("psycopg2.connect") as connect,
            ):
                PostgresSource(
                    host="h",
                    port=5432,
                    database="d",
                    user="u",
                    password="p",
                    table="t",
                    snapshot=bad,
                )
            connect.assert_not_called()

    def test_watermark_requires_value_before_connect(self):
        with (
            self.assertRaisesRegex(ValueError, "snapshot_value"),
            patch("psycopg2.connect") as connect,
        ):
            PostgresSource(
                host="h",
                port=5432,
                database="d",
                user="u",
                password="p",
                table="t",
                snapshot_column="updated_at",
            )
        connect.assert_not_called()

    def test_postgres_watermark_predicate_is_bound(self):
        src = PostgresSource.__new__(PostgresSource)
        src.schema_name, src.table, src.columns = "s", "t", ["a"]
        src.where, src.snapshot_column, src._params = None, "updated_at", ("cut1",)
        self.assertEqual(src._select_sql(), 'SELECT "a" FROM "s"."t" WHERE "updated_at" <= %s')
        self.assertEqual(src._params, ("cut1",))

    def test_snowflake_watermark_combines_with_scope(self):
        src = SnowflakeSource.__new__(SnowflakeSource)
        src.target = MagicMock()
        src.target._qualified.return_value = '"SANTANDER"."PUBLIC"."T"'
        src.table, src.columns = "T", ["a"]
        src.where, src.snapshot_column, src._params = "x = 1", '"TS"', ("cut1",)
        sql = src._select()
        self.assertIn('(x = 1) AND "TS" <= %s', sql)
        self.assertEqual(src._params, ("cut1",))

    def test_exported_snapshot_pins_holder_and_source(self):
        lib = _library({**self.ENV, "source": {**self.ENV["source"], "snapshot": "exported"}})
        holder = MagicMock()
        holder.cursor.return_value.fetchone.return_value = ("0000000A-00000C7B-1",)
        with patch("psycopg2.connect", return_value=holder) as connect:
            lib._bind_snapshot()
        holder.set_session.assert_called_once_with(
            readonly=True, autocommit=False, isolation_level="REPEATABLE READ"
        )
        self.assertEqual(connect.call_count, 1)
        self.assertEqual(lib.env["source"]["snapshot"], "0000000A-00000C7B-1")
        self.assertTrue(lib.meta["snapshot_enforced"])
        self.assertIs(lib._snapshot_holder, holder)

    def test_exported_snapshot_is_postgres_only(self):
        lib = _library({**self.ENV, "source": {"type": "snowflake", "snapshot": "exported"}})
        with self.assertRaisesRegex(ValueError, "postgres-only"):
            lib._bind_snapshot()

    def test_watermark_binding_uses_snapshot_id(self):
        lib = _library(
            {
                **self.ENV,
                "source": {**self.ENV["source"], "snapshot_column": "updated_at"},
            }
        )
        lib._bind_snapshot()
        self.assertEqual(lib.env["source"]["snapshot_value"], "cut-2026-01-01")
        self.assertEqual(lib.meta["source_snapshot"], "watermark:updated_at")
        self.assertTrue(lib.meta["snapshot_enforced"])

    def test_watermark_without_snapshot_id_fails(self):
        lib = _library(
            {
                **self.ENV,
                "run": {},
                "source": {**self.ENV["source"], "snapshot_column": "updated_at"},
            }
        )
        with self.assertRaisesRegex(ValueError, "snapshot_id"):
            lib._bind_snapshot()

    def test_real_target_refuses_unenforced_snapshot(self):
        lib = _library({**self.ENV, "target": {"type": "snowflake", "fake": False}})
        with self.assertRaisesRegex(ValueError, "enforceable snapshot"):
            lib._bind_snapshot()

    def test_emulated_run_records_unenforced_snapshot(self):
        lib = _library({**self.ENV, "source": dict(self.ENV["source"])})
        lib._bind_snapshot()
        self.assertFalse(lib.meta["snapshot_enforced"])
        self.assertEqual(lib.meta["snapshot_id"], "cut-2026-01-01")

    def _pg(self, **attrs):
        src = PostgresSource.__new__(PostgresSource)
        src.schema_name, src.table, src.columns = "s", "t", ["a"]
        src.where, src.snapshot_column, src._params, src._pinned = (
            None,
            None,
            None,
            False,
        )
        src.conn = MagicMock()
        src.conn.cursor.return_value.fetchone.return_value = (7,)
        for key, value in attrs.items():
            setattr(src, key, value)
        return src

    def test_post_snapshot_count_uses_strictly_greater_bound(self):
        src = self._pg(snapshot_column="updated_at", _params=("cut1",))
        self.assertEqual(src.post_snapshot_count(), 7)
        sql, params = src.conn.cursor.return_value.execute.call_args[0]
        self.assertIn('"updated_at" > %s', sql)
        self.assertEqual(params, ("cut1",))

    def test_post_snapshot_count_scoped_by_where(self):
        src = self._pg(where="status = 'active'", snapshot_column="updated_at", _params=("cut1",))
        src.post_snapshot_count()
        sql, _ = src.conn.cursor.return_value.execute.call_args[0]
        self.assertIn("(status = 'active') AND", sql)

    def test_pinned_session_cannot_count_arrivals(self):
        self.assertIsNone(
            self._pg(snapshot_column="u", _params=("c",), _pinned=True).post_snapshot_count()
        )
        self.assertIsNone(self._pg().post_snapshot_count())

    def test_snowflake_post_snapshot_count(self):
        src = SnowflakeSource.__new__(SnowflakeSource)
        src.target = MagicMock()
        src.target._qualified.return_value = '"P"."T"'
        src.target.conn.cursor.return_value.fetchone.return_value = (3,)
        src.table, src.columns = "T", ["a"]
        src.where, src.snapshot_column, src._params = None, '"TS"', ("cut1",)
        self.assertEqual(src.post_snapshot_count(), 3)
        sql, params = src.target.conn.cursor.return_value.execute.call_args[0]
        self.assertIn('"TS" > %s', sql)
        self.assertEqual(params, ("cut1",))

    def test_cutover_block_in_run_summary(self):
        with tempfile.TemporaryDirectory() as tmp:
            results = [
                {
                    "mapping": "a",
                    "stage": "recon",
                    "status": "PASS",
                    "post_snapshot_arrivals": 4,
                    "target_asof": "2026-01-01",
                }
            ]
            meta = {
                "snapshot_enforced": True,
                "source_snapshot": "watermark:u",
                "snapshot_id": "c1",
            }
            write_evidence(results, meta, tmp)
            cutover = json.loads((Path(tmp) / "run_summary.json").read_text())["cutover"]
            self.assertTrue(cutover["snapshot_enforced"])
            self.assertEqual(cutover["post_snapshot_arrivals"], {"a": 4})
            self.assertEqual(cutover["target_asof"], {"a": "2026-01-01"})


class ControlsTests(unittest.TestCase):
    def setUp(self):
        from libs.engine.schema import contract_from_dict

        self.contract = contract_from_dict(
            {
                "contract_name": "controls_test",
                "version": "1",
                "source": {
                    "type": "csv",
                    "path": "x.csv",
                    "columns": ["account_id", "customer_id", "branch_id", "balance", "status"],
                },
                "target": {"type": "postgres", "table": "T"},
                "keys": ["account_id"],
                "columns": [
                    {"name": "account_id", "type": "integer", "nullable": False},
                    {"name": "customer_id", "type": "integer", "nullable": False},
                    {"name": "branch_id", "type": "integer", "nullable": False},
                    {"name": "balance", "type": "decimal", "nullable": True},
                    {"name": "status", "type": "string", "nullable": False},
                ],
                "controls": [
                    {"id": "account_rows", "type": "count"},
                    {"id": "balance_total", "type": "sum", "column": "balance"},
                    {
                        "id": "balance_by_branch",
                        "type": "sum",
                        "column": "balance",
                        "group_by": ["branch_id"],
                    },
                    {"id": "distinct_customers", "type": "count_distinct", "column": "customer_id"},
                ],
            }
        )

    def _frame(self, rows):
        return pd.DataFrame(
            rows, columns=["account_id", "customer_id", "branch_id", "balance", "status"]
        )

    def test_workbook_controls_compile_and_validate(self):
        tables, codes = parser.load_workbook(str(WORKBOOK))
        t = next(t for t in tables if t.mapping_id == "dim_customer_migration")
        contract = builder.build_contract(t, codes, {"type": "postgres"}, {"type": "postgres"})
        specs = engine_controls.validate_controls(contract)
        self.assertEqual({s["id"] for s in specs}, {"dim_customer_rows", "dim_customer_by_type"})
        self.assertEqual(
            next(s for s in specs if s["id"] == "dim_customer_by_type")["group_by"],
            ["customer_type_cd"],
        )

    def test_validate_controls_fails_closed(self):
        contract = copy.deepcopy(self.contract)
        for bad in (
            {"id": "x", "type": "median"},
            {"id": "x", "type": "sum"},  # sum needs a column
            {"id": "x", "type": "count", "column": "balance"},
            {"id": "x", "type": "sum", "column": "nope"},
            {"id": "x", "type": "sum", "column": "balance", "group_by": ["nope"]},
        ):
            with self.subTest(spec=bad):
                contract.raw["controls"] = [bad]
                with self.assertRaises(ValueError):
                    engine_controls.validate_controls(contract)
        contract.raw["controls"] = [
            {"id": "dup", "type": "count"},
            {"id": "dup", "type": "count"},
        ]
        with self.assertRaises(ValueError):
            engine_controls.validate_controls(contract)

    def test_compute_controls_streamed(self):
        from decimal import Decimal

        frames = [
            self._frame([(1, 10, 101, "100.50", "a"), (2, 11, 101, "200.25", "a")]),
            self._frame([(3, 10, 102, "50.00", "a")]),
        ]
        agg = engine_controls.compute_controls(frames, self.contract)
        self.assertEqual(agg["account_rows"][()], 3)
        self.assertEqual(agg["balance_total"][()], Decimal("350.75"))
        self.assertEqual(agg["balance_by_branch"][("101",)], Decimal("300.75"))
        self.assertEqual(agg["balance_by_branch"][("102",)], Decimal("50.00"))
        self.assertEqual(agg["distinct_customers"][()], 2)

    def test_compare_controls_detects_drift_and_missing_groups(self):
        from decimal import Decimal

        expected = {"c": {(): 2, ("b",): Decimal("10")}}
        actual = {"c": {(): 3}}
        diffs = engine_controls.compare_controls(expected, actual)
        self.assertEqual(len(diffs), 2)
        self.assertEqual(
            {(d["control"], d["group"]) for d in diffs},
            {("c", "(all rows)"), ("c", "b")},
        )

    def _trailer(self, tmp, rows):
        p = Path(tmp) / "trailer.csv"
        p.write_text("name,group,value\n" + "\n".join(",".join(r) for r in rows))
        return engine_controls.load_control_file(p)

    def test_trailer_certifies_aggregates(self):
        from decimal import Decimal

        contract = copy.deepcopy(self.contract)
        contract.raw["controls"] = [
            {"id": "t", "type": "sum", "column": "balance", "file_name": "bal"},
            {"id": "g", "type": "count", "group_by": ["status"], "file_name": "by_status"},
        ]
        agg = {"t": {(): Decimal("12.50")}, "g": {("a",): 3}}
        with tempfile.TemporaryDirectory() as tmp:
            declared = self._trailer(tmp, [("bal", "", "12.5"), ("by_status", "a", "3")])
            self.assertEqual(engine_controls.validate_control_file(agg, declared, contract), [])
            declared = self._trailer(tmp, [("bal", "", "99"), ("by_status", "a", "3")])
            diffs = engine_controls.validate_control_file(agg, declared, contract)
            self.assertEqual(len(diffs), 1)
            self.assertEqual(diffs[0]["control"], "t")

    def test_trailer_fails_closed_on_uncertified_control(self):
        contract = copy.deepcopy(self.contract)
        contract.raw["controls"] = [{"id": "t", "type": "count", "file_name": "rows"}]
        with tempfile.TemporaryDirectory() as tmp:
            declared = self._trailer(tmp, [("other", "", "1")])
            diffs = engine_controls.validate_control_file({"t": {(): 5}}, declared, contract)
            self.assertIn("absent from control file", diffs[0]["actual"])
        with self.assertRaises(FileNotFoundError):
            engine_controls.load_control_file("data/trailer/does_not_exist.csv")


class RelationsTests(unittest.TestCase):
    def setUp(self):
        tables, _ = parser.load_workbook(str(WORKBOOK))
        self.rels = {t.mapping_id: t.relations for t in tables if t.relations}

    def test_workbook_relations_parse_and_attach_to_children(self):
        self.assertEqual(set(self.rels), {"dim_customer_address_migration"})
        self.assertEqual(
            self.rels["dim_customer_address_migration"][0],
            {
                "parent_mapping": "dim_customer_migration",
                "parent_column": "involved_party_key",
                "child_column": "involved_party_key",
            },
        )

    def test_key_set_canonicalizes_types(self):
        chunks = [pd.DataFrame({"k": [1, 2, None]}), pd.DataFrame({"k": [" 3 ", 2]})]
        self.assertEqual(engine_relations.key_set(chunks, "k"), {"1", "2", "3"})
        with self.assertRaises(ValueError):
            engine_relations.key_set(chunks, "missing")

    def test_find_orphans_exact_counts_and_nulls_apart(self):
        chunks = [pd.DataFrame({"fk": [1, 9, None, "8", 9]})]
        result = engine_relations.find_orphans(chunks, "fk", {"1", "2"})
        self.assertEqual(result["orphans"], 3)
        self.assertEqual(result["null_fk"], 1)
        self.assertEqual(sorted(result["samples"]), ["8", "9", "9"])

    def test_orphans_are_exact_not_bounded(self):
        chunks = [pd.DataFrame({"fk": list(range(100, 200))})]
        result = engine_relations.find_orphans(chunks, "fk", set(), sample_limit=5)
        self.assertEqual(result["orphans"], 100)
        self.assertEqual(len(result["samples"]), 5)


class NewDialectAdapterTests(unittest.TestCase):
    """mysql/oracle/mssql adapters — SQL generation, quoting, guardrails.
    Driver imports are lazy so these run without the optional
    requirements-adapters.txt drivers installed."""

    def _src(self, cls, dialect):
        src = cls.__new__(cls)
        src.dialect = dialect
        src.schema_name, src.table, src.columns = "s", "t", ["a"]
        src.where, src.snapshot_column, src._params = None, "updated_at", ("cut1",)
        return src

    def test_ident_quoting_per_dialect(self):
        from libs.engine.schema import ident

        self.assertEqual(ident("col", "mysql"), "`col`")
        self.assertEqual(ident("a`b", "mysql"), "`a``b`")
        self.assertEqual(ident("col", "mssql"), "[col]")
        self.assertEqual(ident("a]b", "mssql"), "[a]]b]")
        self.assertEqual(ident("col", "oracle"), '"COL"')
        self.assertEqual(ident("col", "postgres"), '"col"')
        self.assertEqual(ident("col", "snowflake"), '"col"')

    def test_ora_dtype_normalization(self):
        from libs.engine.schema import _ora_dtype

        self.assertEqual(_ora_dtype("NUMBER", 19, 0), "NUMBER_INT")
        self.assertEqual(_ora_dtype("NUMBER", 19, 2), "DECIMAL")
        self.assertEqual(_ora_dtype("NUMBER", None, None), "NUMBER_INT")
        self.assertEqual(_ora_dtype("VARCHAR2", None, None), "VARCHAR2")

    def test_type_maps_cover_all_dialects(self):
        from libs.engine.schema import TYPE_ALIASES, ddl_for_contract

        for dialect in ("postgres", "snowflake", "mysql", "oracle", "mssql"):
            self.assertIn(dialect, TYPE_ALIASES)
            for t in ("integer", "decimal", "string", "date", "timestamp", "boolean"):
                self.assertIn(t, TYPE_ALIASES[dialect])
        for dialect in ("mysql", "oracle", "mssql"):
            ddl = ddl_for_contract(
                type(
                    "C",
                    (),
                    {
                        "columns": [
                            {"name": "id", "type": "integer", "nullable": False},
                            {"name": "amt", "type": "decimal", "nullable": True, "scale": 2},
                            {"name": "nm", "type": "string", "nullable": True, "max_length": 9},
                        ],
                        "keys": ["id"],
                        "schema": "s",
                        "table": "t",
                    },
                )(),
                dialect,
            )
            self.assertIn("PRIMARY KEY", ddl)

    def test_select_sql_uses_dialect_paramstyle(self):
        from libs.adapters.mssql import MssqlSource
        from libs.adapters.mysql import MysqlSource
        from libs.adapters.oracle import OracleSource

        self.assertEqual(
            self._src(MysqlSource, "mysql")._select_sql(),
            "SELECT `a` FROM `s`.`t` WHERE `updated_at` <= %s",
        )
        self.assertEqual(
            self._src(OracleSource, "oracle")._select_sql(),
            'SELECT "A" FROM "S"."T" WHERE "UPDATED_AT" <= :1',
        )
        self.assertEqual(
            self._src(MssqlSource, "mssql")._select_sql(),
            "SELECT [a] FROM [s].[t] WHERE [updated_at] <= ?",
        )

    def test_exported_snapshot_rejected_before_connect(self):
        from libs.adapters.mssql import MssqlSource
        from libs.adapters.mysql import MysqlSource
        from libs.adapters.oracle import OracleSource

        for cls in (MysqlSource, OracleSource, MssqlSource):
            with (
                self.subTest(cls=cls.__name__),
                self.assertRaisesRegex(ValueError, "postgres-only"),
                patch.object(cls, "_connect") as connect,
            ):
                cls(
                    host="h",
                    port=1,
                    database="d",
                    user="u",
                    password="p",
                    table="t",
                    snapshot="0001-0001-1",
                )
            connect.assert_not_called()

    def test_query_guard_accepts_select_only(self):
        from libs.adapters.dbapi import DbapiTarget, _vet_select

        for dialect in ("mysql", "oracle", "mssql"):
            _vet_select("SELECT a FROM t WHERE b = 1", dialect)
            for bad in ("DELETE FROM t", "SELECT 1; DROP TABLE t", "INSERT INTO t VALUES (1)"):
                with self.subTest(dialect=dialect, bad=bad), self.assertRaises(ValueError):
                    _vet_select(bad, dialect)
        tgt = DbapiTarget.__new__(DbapiTarget)
        tgt.dialect, tgt._closed = "mysql", False
        cur = MagicMock()
        cur.description = [("a",)]
        cur.fetchall.return_value = [(1,)]
        tgt.conn = MagicMock()
        tgt.conn.cursor.return_value = cur
        tgt.query("SELECT a FROM t", (1,))
        cur.execute.assert_called_once_with("SELECT a FROM t", (1,))
        with self.assertRaises(ValueError):
            tgt.query("UPDATE t SET a = 1")

    def test_registry_resolves_new_types(self):
        from libs.adapters.registry import _SOURCE_REGISTRY, _TARGET_REGISTRY, _ensure_builtin

        _ensure_builtin()
        for kind in ("mysql", "oracle", "mssql"):
            self.assertIn(kind, _SOURCE_REGISTRY)
            self.assertIn(kind, _TARGET_REGISTRY)

    def test_catalog_dialect_queries(self):
        from libs.adapters.catalog import inspect_schema

        cur = MagicMock()
        cur.fetchall.return_value = []
        conn = MagicMock()
        conn.cursor.return_value = cur
        inspect_schema(conn, "SCH", dialect="oracle")
        self.assertIn("all_tab_columns", cur.execute.call_args[0][0])
        self.assertEqual(cur.execute.call_args[0][1], ("SCH",))
        inspect_schema(conn, "sch", dialect="mssql")
        self.assertIn("information_schema.columns", cur.execute.call_args[0][0])
        self.assertIn("table_schema = ?", cur.execute.call_args[0][0])


class RulesScopeCheckTests(unittest.TestCase):
    """`Check Rules Apply To Environment` — skip only on total scope mismatch."""

    def _lib(self, tables, present, schema="wdr_cl_fds_core"):
        from types import SimpleNamespace

        lib = ReconciliationLibrary()
        lib.env = {"target": {"schema": schema}}
        lib.excel_rules = [SimpleNamespace(enabled=True, target_table=t) for t in tables]
        target = MagicMock()
        target.query.return_value = pd.DataFrame({"table_name": list(present)})
        lib._ensure_target = MagicMock(return_value=target)
        return lib

    def test_all_rule_tables_missing_returns_skip_reason(self):
        lib = self._lib(["cup00301", "cup01201", "v_dim_customer_src"], present=[])
        reason = lib.check_rules_apply()
        self.assertIn("different hop", reason)
        self.assertIn("wdr_cl_fds_core", reason)

    def test_all_rule_tables_present_allows_run(self):
        lib = self._lib(
            ["cup00301", "v_dim_customer_src"],
            present=["cup00301", "v_dim_customer_src"],
            schema="signature",
        )
        self.assertEqual(lib.check_rules_apply(), "")

    def test_partial_presence_allows_run_so_broken_rules_error(self):
        # One table exists, one doesn't → no blanket skip; the missing-table
        # rule must surface its own ERROR, never a silent pass.
        lib = self._lib(
            ["cup00301", "typo_table"],
            present=["cup00301"],
            schema="signature",
        )
        self.assertEqual(lib.check_rules_apply(), "")

    def test_no_rules_or_catalog_failure_returns_empty(self):
        lib = self._lib([], present=[])
        self.assertEqual(lib.check_rules_apply(), "")
        lib = self._lib(["cup00301"], present=[])
        lib._ensure_target.return_value.query.side_effect = RuntimeError("no catalog")
        self.assertEqual(lib.check_rules_apply(), "")


if __name__ == "__main__":
    unittest.main()
