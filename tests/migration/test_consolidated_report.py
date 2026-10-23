"""Unit tests for the consolidated end-to-end lifecycle report merger."""

import tempfile
import unittest
from pathlib import Path

import openpyxl

from libs.excelrules.models import Rule, RuleResult, Violation
from libs.reporting import consolidated_report
from libs.reporting.review_pdf import build_review, write_review_pdf


def _run(verdict="PASS"):
    return {
        "verdict": verdict,
        "coverage": "COMPLETE",
        "cutover": {
            "snapshot_id": "2026-10-15",
            "source_snapshot": "00000003-00000164-1",
            "snapshot_enforced": True,
        },
        "results": [
            {
                "mapping": "csv_cup00301",
                "stage": "recon",
                "status": "PASS",
                "detail": "200 source rows, 200 target rows",
                "source_count": 200,
                "target_count": 200,
                "seconds": 1.2,
            },
            {
                "mapping": "csv_cup00301",
                "stage": "controls",
                "status": "FAIL" if verdict == "FAIL" else "PASS",
                "detail": "8 control(s) certified",
                "control_mismatches": (
                    [{"control": "naics_sum", "group": "", "expected": "1", "actual": "2"}]
                    if verdict == "FAIL"
                    else []
                ),
            },
        ],
    }


def _results(status_violations):
    rule = Rule(
        test_id="BR-AGG-UNIQ-CUNBR",
        rule_type="CUSTOM_SQL",
        target_table="cup00301",
        target_column="cunbr",
        severity="CRITICAL",
        key_column="cunbr",
        business_rule="customer number must be unique",
    )
    return [
        RuleResult(
            rule=rule,
            violations=status_violations,
            sql="SELECT cunbr FROM signature.cup00301 WHERE ...",
        )
    ]


class ConsolidatedReportTest(unittest.TestCase):
    def _write(self, verdict="PASS", violations=None):
        tmp = tempfile.mkdtemp()
        path = consolidated_report.write_consolidated(
            _run(verdict),
            _results(violations or []),
            tmp,
            meta={"environment": "phase1-hop1", "timestamp": "2026-10-15T00:00:00"},
        )
        return Path(tmp), path

    def test_writes_workbook_and_markdown(self):
        tmp, path = self._write()
        self.assertTrue(path.exists())
        self.assertTrue((tmp / "end_to_end_report.md").exists())
        self.assertTrue((tmp / "end_to_end_report.pdf").read_bytes().startswith(b"%PDF-"))
        wb = openpyxl.load_workbook(path)
        self.assertEqual(
            sorted(wb.sheetnames),
            [
                "Column Profile",
                "Controls",
                "Data Diffs",
                "Insights",
                "Metadata",
                "Queries",
                "Rules",
                "Run Info",
                "Stages",
                "Violations",
            ],
        )

    def test_overall_status_reflects_failures(self):
        tmp, path = self._write(verdict="FAIL")
        wb = openpyxl.load_workbook(path)
        vals = {r[0].value: r[1].value for r in wb["Run Info"].iter_rows(min_row=2)}
        self.assertEqual(vals["Overall Status"], "FAIL")
        self.assertEqual(vals["Verdict"], "FAIL")
        # control mismatch is flattened onto the Controls sheet
        rows = list(wb["Controls"].iter_rows(min_row=2, values_only=True))
        self.assertIn(("csv_cup00301", "naics_sum", None, "1", "2"), rows)

    def test_rules_sheet_carries_plain_english(self):
        tmp, path = self._write()
        wb = openpyxl.load_workbook(path)
        rows = list(wb["Rules"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(rows[0][0], "BR-AGG-UNIQ-CUNBR")
        self.assertEqual(rows[0][2], "customer number must be unique")
        self.assertEqual(rows[0][4], "PASS")

    def test_violations_flattened(self):
        viol = [Violation("BR-1", "cunbr", "expected-9-digit", "12345", key="C000001")]
        tmp, path = self._write(violations=viol)
        wb = openpyxl.load_workbook(path)
        rows = list(wb["Violations"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(rows[0][0], "BR-AGG-UNIQ-CUNBR")
        self.assertEqual(rows[0][4], "expected-9-digit")
        md = (tmp / "end_to_end_report.md").read_text(encoding="utf-8")
        self.assertIn("BR-AGG-UNIQ-CUNBR", md)


class ReviewPdfTest(unittest.TestCase):
    def setUp(self):
        self.run = _run()
        self.run["meta"] = {
            "selected": ["csv_cup00301"],
            "relation_children": [],
            "environment": "synthetic-test",
            "workbook_sha256": "a" * 64,
        }
        self.run["not_run"] = {"csv_cup00301": []}
        self.run["results"] += [
            {"mapping": "csv_cup00301", "stage": s, "status": "PASS"}
            for s in ("mapping", "source", "dq", "schema", "transform", "profile")
        ]
        self.run["results"][0].update(missing=0, extra=0, diffs=0)
        self.outcomes = [
            {"test": "Migration.01 End To End Migration Testing." + name, "status": "PASS"}
            for name in (
                "Requirements Intake And Mapping Compile",
                "Source And Schema Conformance",
                "Transformation And Record Reconciliation",
                "Business Validation Controls Relations Profile And Rules",
                "Lifecycle Evidence Is Complete And Passing",
            )
        ]

    def review(self, results=None, **kwargs):
        return build_review(self.run, results or [], test_outcomes=self.outcomes, **kwargs)

    def test_complete_run_without_additional_rules(self):
        review = self.review(declared_rules=[])
        self.assertEqual(review["status"], "PASS")
        self.assertEqual(len(review["cases"]), 5)
        self.assertEqual(review["rules"], [])
        self.assertEqual(review["datasets"][0]["counts"], [200, 200, 0, 0, 0])

    def test_failed_rule_overrides_successful_stage_and_test_results(self):
        results = _results([Violation("BR-1", "cunbr", "PRIVATE-EXPECTED", "PRIVATE-ACTUAL")])
        review = self.review(results)
        self.assertEqual(review["status"], "FAIL")
        self.assertEqual(review["rules"][0]["evidence"], "1 reported violation(s).")
        self.assertNotIn("PRIVATE", str(review))
        self.assertNotIn("SELECT", str(review))

    def test_exact_counts_are_reported_without_compared_values(self):
        self.run["results"][0].update(status="FAIL", source_count=195, extra=5)
        self.run["results"][1].update(
            status="FAIL", control_mismatches=[{"actual": "PRIVATE-SAMPLE"}] * 8
        )
        review = self.review()
        self.assertEqual(review["datasets"][0]["counts"], [195, 200, 0, 5, 0])
        self.assertIn(["Certified totals", "1 FAIL; 8 total mismatches"], review["checks"])
        self.assertNotIn("PRIVATE", str(review))
        del self.run["results"][0]["target_count"]
        self.assertEqual(self.review()["datasets"][0]["counts"][1], "Not recorded")

    def test_missing_rule_without_requirement_id_is_incomplete(self):
        review = self.review(declared_rules=[_results([])[0].rule])
        self.assertEqual(review["status"], "INCOMPLETE")
        self.assertEqual(review["rules"][0]["status"], "NOT RUN")

    def test_partial_evidence_and_missing_test_outcomes_never_pass(self):
        self.run["results"] = [r for r in self.run["results"] if r["stage"] != "transform"]
        self.assertEqual(self.review()["status"], "INCOMPLETE")
        self.assertEqual(build_review(self.run, [])["status"], "INCOMPLETE")

    def test_skipped_and_failed_test_outcomes_are_respected(self):
        for raw, expected in (("SKIP", "INCOMPLETE"), ("FAIL", "FAIL"), ("UNKNOWN", "INCOMPLETE")):
            with self.subTest(status=raw):
                self.outcomes[-1]["status"] = raw
                self.assertEqual(self.review()["status"], expected)

    def test_warnings_require_review_and_raw_errors_are_not_shared(self):
        result = _results([Violation("BR-1", "cunbr", "private", "private")])[0]
        result.rule.severity = "LOW"
        self.assertEqual(self.review([result])["status"], "REVIEW")
        result.error = "PRIVATE connection error"
        self.run["results"][0]["detail"] = "PRIVATE row contents"
        review = self.review([result])
        self.assertEqual(review["status"], "FAIL")
        self.assertNotIn("PRIVATE", str(review))

    def test_pdf_handles_long_rule_text_and_multiple_pages(self):
        results = []
        for index in range(35):
            result = _results([])[0]
            result.rule.test_id = f"BR-{index}"
            result.rule.business_rule = "Check <value> & its declared format. " * 30
            results.append(result)
        with tempfile.TemporaryDirectory() as tmp:
            path = write_review_pdf(self.run, results, tmp, test_outcomes=self.outcomes)
            data = path.read_bytes()
            self.assertTrue(data.startswith(b"%PDF-"))
            self.assertIn(b"%%EOF", data[-100:])
            self.assertGreater(data.count(b"/Type /Page\n"), 1)


if __name__ == "__main__":
    unittest.main()
