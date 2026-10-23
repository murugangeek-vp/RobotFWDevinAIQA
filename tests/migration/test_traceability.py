"""Unit tests for RTM / DC-DA-DI-DG scorecard / defect-report builders."""

import tempfile
import unittest
from pathlib import Path

import openpyxl

from libs.excelrules.models import Rule, RuleResult, Violation
from libs.mapping.models import FieldMapping, TableMapping
from libs.reporting import traceability


def _table():
    t = TableMapping(
        mapping_id="csv_cup00301",
        source_schema="csv",
        source_table="cup00301",
        target_schema="signature",
        target_table="cup00301",
        keys=["cunbr"],
        requirement_id="REQ-DC-101",
    )
    t.fields = [
        FieldMapping(
            mapping_id="csv_cup00301", source_column="cunbr", target_column="cunbr",
            target_type="string", pii=True, requirement_id="REQ-DC-101",
        )
    ]
    t.controls = [{"id": "c1", "type": "count", "requirement_id": "REQ-DC-101"}]
    t.relations = [
        {"parent_mapping": "csv_cup00301", "parent_column": "cunbr",
         "child_column": "cunbr", "requirement_id": "REQ-DI-102"}
    ]
    return t


def _rule():
    return Rule(
        test_id="BR-1", rule_type="CUSTOM_SQL", target_table="cup00301",
        target_column="cunbr", requirement_id="REQ-DI-103",
    )


def _evidence(status="PASS"):
    return [
        {"mapping": "csv_cup00301", "stage": "recon", "status": status,
         "detail": "200/200", "control_mismatches": []},
        {"mapping": "csv_cup00301", "stage": "controls", "status": status,
         "detail": "", "control_mismatches":
            [{"control": "c1", "group": "", "expected": "1", "actual": "2"}]
            if status == "FAIL" else []},
        {"mapping": "csv_cup00301", "stage": "schema", "status": status, "detail": ""},
        {"mapping": "csv_cup00301", "stage": "relations", "status": status, "detail": ""},
        {"mapping": "csv_cup00301", "stage": "transform", "status": status, "detail": ""},
    ]


class RtmTest(unittest.TestCase):
    def test_requirement_links_stage_checks(self):
        rows, uncovered = traceability.build_rtm(
            [_table()], [_rule()], [], _evidence()
        )
        reqs = {r["req"] for r in rows}
        self.assertIn("REQ-DC-101", reqs)
        self.assertIn("REQ-DI-102", reqs)
        self.assertIn("REQ-DI-103", reqs)
        self.assertEqual(uncovered, ["REQ-DI-103"])  # rule never ran -> UNCOVERED

    def test_control_mismatch_fails_its_requirement(self):
        rows, _ = traceability.build_rtm([_table()], [], [], _evidence("FAIL"))
        c1 = [r for r in rows if r["check"] == "certified control"]
        self.assertTrue(all(r["status"] == "FAIL" for r in c1))

    def test_rule_req_reports_result_status(self):
        res = [RuleResult(rule=_rule(), violations=[
            Violation("BR-1", "cunbr", "e", "a", key="K1")
        ])]
        rows, uncovered = traceability.build_rtm(
            [_table()], [_rule()], res, _evidence()
        )
        rule_row = next(r for r in rows if r["check"] == "business rule")
        self.assertEqual(rule_row["status"], "FAIL")
        self.assertEqual(uncovered, [])


class ScorecardTest(unittest.TestCase):
    def test_dimensions_and_verdicts(self):
        run = {"results": _evidence(), "cutover": {"snapshot_enforced": True},
               "meta": {"started": "x"}}
        rows, verdicts = traceability.build_scorecard(
            [_table()], [], run["results"], run
        )
        self.assertEqual(set(verdicts), {"DC", "DA", "DI", "DG"})
        self.assertEqual(verdicts["DC"], "CONFIRMED")
        self.assertEqual(verdicts["DG"], "CONFIRMED")

    def test_failing_stage_fails_its_dimension(self):
        run = {"results": _evidence("FAIL"), "cutover": {}, "meta": {}}
        _, verdicts = traceability.build_scorecard(
            [_table()], [], run["results"], run
        )
        self.assertEqual(verdicts["DC"], "FAILED")


class DefectsAndFilesTest(unittest.TestCase):
    def test_defect_report_only_on_failure(self):
        tmp = Path(tempfile.mkdtemp())
        ok = traceability.write_trace(
            [_table()], [_rule()], [], {"results": _evidence()}, tmp
        )
        self.assertIsNone(ok["files"]["defects"])
        self.assertFalse((tmp / "defect_report.md").exists())

        bad = traceability.write_trace(
            [_table()], [_rule()], [], {"results": _evidence("FAIL")}, tmp
        )
        self.assertTrue((tmp / "defect_report.md").exists())
        md = (tmp / "defect_report.md").read_text(encoding="utf-8")
        self.assertIn("csv_cup00301/controls", md)
        self.assertIn("REQ-DC-101", md)

    def test_artifacts_written(self):
        tmp = Path(tempfile.mkdtemp())
        traceability.write_trace(
            [_table()], [_rule()], [], {"results": _evidence()}, tmp
        )
        wb = openpyxl.load_workbook(tmp / "rtm.xlsx")
        self.assertEqual(wb.active.title, "RTM")
        wb2 = openpyxl.load_workbook(tmp / "compliance_scorecard.xlsx")
        self.assertIn("Dimension Verdicts", wb2.sheetnames)


if __name__ == "__main__":
    unittest.main()
